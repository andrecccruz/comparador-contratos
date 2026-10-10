# -*- coding: utf-8 -*-
"""
recursos.py
Funcoes da Fase 5 sem dependencia do Flask (testaveis isoladamente):
rodadas de negociacao, biblioteca de clausulas aprovadas, painel do historico,
comparacao em lote e versoes de modelo.
"""

import os
import re
import io
import csv
import json
import shutil
import threading
from datetime import datetime
from collections import Counter, defaultdict

import comparar_contratos as motor

TIPOS_MUDANCA = ("ALTERADO", "INCLUIDO", "REMOVIDO", "MOVIDO")
_trava_arquivos = threading.RLock()


def ler_json(caminho, padrao=None):
    try:
        with open(caminho, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return padrao


def gravar_json(caminho, dados):
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    tmp = caminho + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=1)
    os.replace(tmp, caminho)


# ----------------------------------------------------------------- rodadas

def chave_fornecedor(execucao):
    """CNPJ quando houver; senao o nome normalizado."""
    cnpj = re.sub(r"\D", "", execucao.get("cnpj") or "")
    if len(cnpj) == 14:
        return "cnpj:" + cnpj
    nome = motor.sem_acento(execucao.get("fornecedor") or "").upper()
    nome = re.sub(r"\b(LTDA|S\.?\s?A\.?|EIRELI|ME|EPP)\b\.?", "", nome)
    return "nome:" + re.sub(r"[^A-Z0-9]+", "", nome)


def rodadas_do_grupo(execucoes, execucao):
    """Execucoes concluidas do mesmo fornecedor + modelo + departamento, em ordem."""
    alvo = (execucao.get("departamento"), chave_fornecedor(execucao), execucao.get("modelo"))
    grupo = [e for e in execucoes.values()
             if e.get("status") == "concluido"
             and (e.get("departamento"), chave_fornecedor(e), e.get("modelo")) == alvo]
    return sorted(grupo, key=lambda e: e.get("criado_em_iso", ""))


def numerar_rodadas(execucoes):
    """Grava 'rodada' (1, 2, 3...) em cada execucao concluida. Idempotente."""
    grupos = defaultdict(list)
    for e in execucoes.values():
        if e.get("status") == "concluido":
            grupos[(e.get("departamento"), chave_fornecedor(e), e.get("modelo"))].append(e)
    for itens in grupos.values():
        itens.sort(key=lambda e: e.get("criado_em_iso", ""))
        for n, e in enumerate(itens, 1):
            e["rodada"] = n
            e["rodadas_total"] = len(itens)


def _resumo_oc(oc):
    return {"id": oc["id"], "tipo": oc["tipo"], "clausula": oc["clausula"],
            "risco": oc.get("risco", 0), "selo": oc.get("selo", ""),
            "texto_modelo": oc.get("texto_modelo", ""),
            "texto_devolucao": oc.get("texto_devolucao", "")}


def comparar_rodadas(anterior, atual):
    """Classifica o que o fornecedor fez de uma rodada para a seguinte.

    aceitas         mudanca da rodada anterior que sumiu (voltou ao texto do modelo)
    mantidas        mesma mudanca nas duas rodadas
    mudou_de_novo   o mesmo trecho do modelo mudou de outro jeito
    novas           mudanca que so aparece na rodada atual
    resolvidas      campo pendente que foi preenchido"""
    def mudancas(res):
        return [o for o in res.get("ocorrencias", []) if o["tipo"] in TIPOS_MUDANCA]

    ant, atu = mudancas(anterior), mudancas(atual)
    chaves_atu = {o["chave"]: o for o in atu}
    por_modelo_atu = defaultdict(list)
    for o in atu:
        if o.get("texto_modelo"):
            por_modelo_atu[motor.normalizar(o["texto_modelo"])].append(o)
    usadas = set()
    saida = {"aceitas": [], "mantidas": [], "mudou_de_novo": [], "novas": [], "resolvidas": []}
    for o in ant:
        if o["chave"] in chaves_atu:
            saida["mantidas"].append(_resumo_oc(chaves_atu[o["chave"]]))
            usadas.add(o["chave"])
            continue
        mesmos = [x for x in por_modelo_atu.get(motor.normalizar(o.get("texto_modelo", "")), [])
                  if x["chave"] not in usadas] if o.get("texto_modelo") else []
        if mesmos:
            novo = mesmos[0]
            usadas.add(novo["chave"])
            item = _resumo_oc(novo)
            item["antes"] = o.get("texto_devolucao", "")
            saida["mudou_de_novo"].append(item)
        else:
            saida["aceitas"].append(_resumo_oc(o))
    for o in atu:
        if o["chave"] not in usadas:
            saida["novas"].append(_resumo_oc(o))
    pend_ant = {motor.normalizar(o["texto_modelo"]) for o in anterior.get("ocorrencias", [])
                if o["tipo"] == "PENDENTE"}
    for o in atual.get("ocorrencias", []):
        if o["tipo"] == "PREENCHIDO" and motor.normalizar(o["texto_modelo"]) in pend_ant:
            saida["resolvidas"].append(_resumo_oc(o))
    saida["contagem"] = {k: len(v) for k, v in saida.items() if isinstance(v, list)}
    return saida


# ----------------------------------------------------------------- biblioteca

def caminho_biblioteca(pasta_dados, departamento):
    return os.path.join(pasta_dados, "_biblioteca", "%s.json" % departamento)


def carregar_biblioteca(pasta_dados, departamento):
    return ler_json(caminho_biblioteca(pasta_dados, departamento), {"itens": {}})


def mapa_biblioteca(pasta_dados, departamento, excluir_execucao=None):
    """{chave_texto: {data, fornecedor}} para o motor marcar 'ja aprovado antes'."""
    itens = carregar_biblioteca(pasta_dados, departamento).get("itens", {})
    return {k: v for k, v in itens.items() if v.get("execucao") != excluir_execucao}


def atualizar_biblioteca(pasta_dados, execucao, resultado, revisao):
    """Guarda os trechos ACEITOS desta execucao (substitui os anteriores dela)."""
    if not resultado:
        return 0
    with _trava_arquivos:
        caminho = caminho_biblioteca(pasta_dados, execucao["departamento"])
        bib = ler_json(caminho, {"itens": {}})
        bib["itens"] = {k: v for k, v in bib.get("itens", {}).items()
                        if v.get("execucao") != execucao["id"]}
        itens = (revisao or {}).get("itens", {})
        novos = 0
        for oc in resultado.get("ocorrencias", []):
            if oc["tipo"] not in ("ALTERADO", "INCLUIDO", "MOVIDO") or not oc["texto_devolucao"]:
                continue
            if itens.get(oc.get("chave"), {}).get("status") != "aceito":
                continue
            chave = motor.chave_texto(oc["texto_devolucao"])
            bib["itens"][chave] = {
                "texto": oc["texto_devolucao"][:4000], "clausula": oc["clausula"],
                "tipo": oc["tipo"], "data": datetime.now().strftime("%d/%m/%Y"),
                "fornecedor": execucao.get("fornecedor", ""), "execucao": execucao["id"],
                "modelo": execucao.get("modelo", ""),
                "comentario": itens[oc["chave"]].get("comentario", "")[:500]}
            novos += 1
        bib["atualizado_em"] = datetime.now().isoformat(timespec="seconds")
        gravar_json(caminho, bib)
        return novos


# ----------------------------------------------------------------- painel

def _data(execucao):
    try:
        return datetime.fromisoformat(execucao.get("criado_em_iso", ""))
    except ValueError:
        return None


def clausulas_alteradas(resultado):
    """Contagem de mudancas reais (sem preenchimento) por clausula."""
    cont = Counter()
    for o in (resultado or {}).get("ocorrencias", []):
        if o["tipo"] in TIPOS_MUDANCA:
            cont[o["clausula"]] += 1
    return dict(cont)


def filtrar_execucoes(execucoes, departamento="", fornecedor="", de="", ate="", status="",
                      busca=""):
    saida = []
    forn = motor.sem_acento(fornecedor).lower().strip()
    termo = motor.sem_acento(busca).lower().strip()
    for e in execucoes.values():
        if departamento and e.get("departamento") != departamento:
            continue
        iso = (e.get("criado_em_iso") or "")[:10]
        if de and iso < de:
            continue
        if ate and iso > ate:
            continue
        if status and e.get("status_revisao", "pendente") != status:
            continue
        if forn and forn not in motor.sem_acento(e.get("fornecedor") or "").lower():
            continue
        if termo:
            alvo = motor.sem_acento(" ".join(str(e.get(k) or "") for k in (
                "arquivo", "fornecedor", "cnpj", "modelo", "responsavel"))).lower()
            if termo not in alvo:
                continue
        saida.append(e)
    return sorted(saida, key=lambda e: e.get("criado_em_iso", ""), reverse=True)


def painel(execucoes_filtradas):
    concluidas = [e for e in execucoes_filtradas if e.get("status") == "concluido"]
    riscos = [e["risco_geral"] for e in concluidas if isinstance(e.get("risco_geral"), int)]
    clausulas = Counter()
    por_fornecedor = defaultdict(list)
    tempos = []
    for e in concluidas:
        clausulas.update(e.get("clausulas_alteradas") or {})
        if isinstance(e.get("risco_geral"), int):
            por_fornecedor[e.get("fornecedor") or "—"].append(e["risco_geral"])
        fim = e.get("revisao_concluida_em_iso")
        ini = _data(e)
        if fim and ini:
            try:
                tempos.append((datetime.fromisoformat(fim) - ini).total_seconds() / 86400.0)
            except ValueError:
                pass
    faixas = [("até 1 dia", 0, 1), ("1 a 3 dias", 1, 3), ("3 a 7 dias", 3, 7),
              ("7 a 15 dias", 7, 15), ("mais de 15 dias", 15, 1e9)]
    risco_fornecedor = sorted(((f, round(sum(v) / len(v)), len(v))
                               for f, v in por_fornecedor.items()), key=lambda x: -x[1])[:10]
    return {
        "risco_fornecedor_barras": [("%s (%d)" % (f, n), media)
                                    for f, media, n in risco_fornecedor],
        "total": len(execucoes_filtradas),
        "concluidas": len(concluidas),
        "risco_medio": round(sum(riscos) / len(riscos)) if riscos else None,
        "incompletas": sum(1 for e in concluidas if e.get("devolucao_incompleta")),
        "fora_politica": sum(1 for e in concluidas if e.get("fora_da_politica")),
        "revisadas": sum(1 for e in concluidas if e.get("status_revisao") in ("aprovado",
                                                                              "devolvido")),
        "por_status": dict(Counter(e.get("status_revisao", "pendente") for e in concluidas)),
        "clausulas": clausulas.most_common(10),
        "risco_fornecedor": risco_fornecedor,
        "tempo_medio_dias": round(sum(tempos) / len(tempos), 1) if tempos else None,
        "tempo_faixas": [(rot, sum(1 for t in tempos if a <= t < b)) for rot, a, b in faixas],
    }


COLUNAS_CSV = ["criado_em", "departamento", "arquivo", "fornecedor", "cnpj", "modelo",
               "rodada", "status", "status_revisao", "responsavel", "risco_geral", "selo_risco",
               "preenchidos", "alterados", "incluidos", "removidos", "movidos", "pendentes",
               "criticos", "fora_da_politica", "devolucao_incompleta", "versao_motor"]


def csv_execucoes(execucoes):
    saida = io.StringIO()
    w = csv.writer(saida, delimiter=";")
    w.writerow(COLUNAS_CSV)
    for e in execucoes:
        w.writerow([("sim" if e.get(c) is True else "" if e.get(c) is False else
                     e.get(c, "")) for c in COLUNAS_CSV])
    return "﻿" + saida.getvalue()


# ----------------------------------------------------------------- lote

def resumo_lote(execucoes, lote_id):
    itens = sorted((e for e in execucoes.values() if e.get("lote") == lote_id),
                   key=lambda e: (-(e.get("risco_geral") if isinstance(e.get("risco_geral"), int)
                                    else -1), e.get("arquivo", "")))
    concluidas = [e for e in itens if e.get("status") == "concluido"]
    clausulas = Counter()
    for e in concluidas:
        clausulas.update(e.get("clausulas_alteradas") or {})
    totais = Counter()
    for e in concluidas:
        for k in ("preenchidos", "alterados", "incluidos", "removidos", "movidos",
                  "pendentes", "criticos"):
            totais[k] += e.get(k) or 0
    riscos = [e["risco_geral"] for e in concluidas if isinstance(e.get("risco_geral"), int)]
    return {"id": lote_id, "itens": itens, "concluidas": len(concluidas),
            "pendentes_fila": sum(1 for e in itens if e.get("status") in ("na_fila",
                                                                          "processando")),
            "erros": sum(1 for e in itens if e.get("status") == "erro"),
            "totais": dict(totais), "clausulas": clausulas.most_common(15),
            "risco_medio": round(sum(riscos) / len(riscos)) if riscos else None,
            "risco_max": max(riscos) if riscos else None,
            "modelo": itens[0].get("modelo") if itens else "",
            "departamento": itens[0].get("departamento") if itens else ""}


def xlsx_lote(resumo, cor_cabecalho="274BC0"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "LOTE"
    cab = ["Arquivo", "Fornecedor", "CNPJ", "Status", "Risco", "Nível", "Preenchidos",
           "Alterados", "Incluídos", "Removidos", "Movidos", "Pendentes", "Críticos",
           "Fora da política", "Incompleta"]
    ws.append(cab)
    for e in resumo["itens"]:
        ws.append([e.get("arquivo"), e.get("fornecedor"), e.get("cnpj"), e.get("status"),
                   e.get("risco_geral"), e.get("selo_risco"), e.get("preenchidos"),
                   e.get("alterados"), e.get("incluidos"), e.get("removidos"),
                   e.get("movidos"), e.get("pendentes"), e.get("criticos"),
                   e.get("fora_da_politica"), "SIM" if e.get("devolucao_incompleta") else ""])
    wc = wb.create_sheet("CLAUSULAS")
    wc.append(["Cláusula", "Mudanças no lote"])
    for c, n in resumo["clausulas"]:
        wc.append([c, n])
    fill = PatternFill("solid", fgColor=cor_cabecalho.lstrip("#"))
    for folha in (ws, wc):
        for cel in folha[1]:
            cel.fill, cel.font = fill, Font(color="FFFFFF", bold=True)
            cel.alignment = Alignment(vertical="top")
        folha.column_dimensions["A"].width = 46
        folha.column_dimensions["B"].width = 34
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ----------------------------------------------------------------- versoes de modelo

RE_VERSAO = re.compile(r"^\d{8}_\d{6}_[\w\-. ]+$")


def pasta_versoes(pasta_modelos_dep, nome):
    base = os.path.splitext(nome)[0]
    return os.path.join(pasta_modelos_dep, "_versoes", base)


def arquivar_versao(pasta_modelos_dep, nome, meta):
    """Move o modelo atual para _versoes/<base>/<carimbo>_<nome> e registra no meta."""
    atual = os.path.join(pasta_modelos_dep, nome)
    if not os.path.isfile(atual):
        return None
    destino_dir = pasta_versoes(pasta_modelos_dep, nome)
    os.makedirs(destino_dir, exist_ok=True)
    carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
    arquivo = "%s_%s" % (carimbo, nome)
    shutil.move(atual, os.path.join(destino_dir, arquivo))
    registro = meta.setdefault(nome, {})
    versoes = registro.setdefault("versoes", [])
    versoes.insert(0, {"arquivo": arquivo, "arquivado_em": datetime.now().strftime(
        "%d/%m/%Y %H:%M"), "data_modelo": registro.get("data", ""),
        "observacao": registro.get("observacao", "")})
    registro.pop("data", None)                  # a data manual era da versao anterior
    return arquivo


def caminho_versao(pasta_modelos_dep, nome, arquivo_versao):
    if not RE_VERSAO.match(arquivo_versao or "") or "/" in arquivo_versao \
            or "\\" in arquivo_versao:
        return None
    pasta = os.path.abspath(pasta_versoes(pasta_modelos_dep, nome))
    caminho = os.path.abspath(os.path.join(pasta, arquivo_versao))
    if not caminho.startswith(pasta + os.sep) or not os.path.isfile(caminho):
        return None
    return caminho
