# -*- coding: utf-8 -*-
"""
relatorios.py
Gera as saidas a partir do resultado estruturado (resultado.json):
  - comparativo.html  interativo, arquivo unico autocontido (ver comparativo_interativo.py)
  - comparativo.pdf   template estatico proprio para xhtml2pdf (sem JS / CSS moderno)
  - alteracoes.xlsx   RESUMO, FICHA, DETALHE, PENDENCIAS, PLAYBOOK
  - resumo_alteracoes.txt
Padrao visual: docs/identidade.md
"""

import os
import html
import logging

import identidade as identidade_mod

TIPOS = ("PREENCHIDO", "ALTERADO", "INCLUIDO", "REMOVIDO", "MOVIDO", "PENDENTE",
         "FORMATACAO")

# cores semanticas (docs/identidade.md, secao 3): texto, fundo
CORES_TIPO = {
    "PREENCHIDO": ("#0E6377", "#E0F2F7"),
    "ALTERADO": ("#8A5F00", "#FFF4D6"),
    "INCLUIDO": ("#1F7A3D", "#D6F5DF"),
    "REMOVIDO": ("#B00020", "#FFDDE1"),
    "MOVIDO": ("#4A4AB5", "#E3E3F7"),
    "PENDENTE": ("#8A4B08", "#FBEBD7"),
    "FORMATACAO": ("#5C6570", "#EEF0F2"),
}
CORES_RISCO = {"Baixo": "#1F7A3D", "Médio": "#8A6100", "Alto": "#B00020"}
NOMES_TIPO = {"PREENCHIDO": "Preenchido", "ALTERADO": "Alterado", "INCLUIDO": "Incluído",
              "REMOVIDO": "Removido", "MOVIDO": "Movido", "PENDENTE": "Pendente",
              "FORMATACAO": "Formatação"}
ROTULOS_BLOCO = {
    "ALTERADO": ("MODELO PADRÃO", "DEVOLUÇÃO DO FORNECEDOR"),
    "REMOVIDO": ("SUPRIMIDO PELO FORNECEDOR", None),
    "INCLUIDO": (None, "INCLUÍDO PELO FORNECEDOR"),
    "MOVIDO": (None, "TEXTO MOVIDO NA DEVOLUÇÃO"),
    "FORMATACAO": ("MODELO PADRÃO", "DEVOLUÇÃO (SÓ FORMATAÇÃO)"),
}


def _ident(resultado, cfg):
    return identidade_mod.mesclar((cfg or {}).get("identidade") or resultado.get("identidade"))


def texto_rodape(resultado, cfg=None):
    ident = _ident(resultado, cfg)
    partes = ["%s · motor v%s" % (ident["nome_produto"],
                                       resultado["meta"].get("versao_motor", ""))]
    if ident.get("rodape_relatorio"):
        partes.append(ident["rodape_relatorio"])
    return " · ".join(partes)


def _fmt_num(valor, casas=0):
    if valor is None:
        return "-"
    texto = ("{:,.%df}" % casas).format(valor)
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def linhas_ficha(resultado):
    """Pares (rotulo, valor, alerta) da ficha, na ordem de exibicao."""
    f = resultado.get("ficha") or {}
    meta = resultado["meta"]
    partes = f.get("partes", {})
    vig = f.get("vigencia", {})
    ind = f.get("indicadores", {})
    cnpj = meta.get("cnpj") or ""
    cnpj_ok = next((c["valido"] for c in f.get("cnpjs", []) if c["numero"] == cnpj), None)
    linhas = [
        ("Contratante", partes.get("contratante") or "-", False),
        ("Contratada", meta.get("fornecedor") or "-", False),
        ("CNPJ", (cnpj + (" — válido" if cnpj_ok else " — INVÁLIDO"
                          if cnpj_ok is False else "")) if cnpj else "-", cnpj_ok is False),
        ("Nº do contrato", f.get("numero_contrato") or "-", False),
    ]
    if vig.get("inicio") or vig.get("meses"):
        texto = "%s a %s" % (vig.get("inicio") or "?", vig.get("fim") or "?") \
            if vig.get("inicio") else ""
        if vig.get("meses"):
            texto += (" (%d meses)" if texto else "%d meses") % vig["meses"]
        linhas.append(("Vigência", texto, vig.get("valida") is False))
    if ind.get("pagamento_dias") is not None:
        linhas.append(("Pagamento", "%d dias" % ind["pagamento_dias"], False))
    if ind.get("multa_percentual_max") is not None:
        linhas.append(("Multa (maior)", "%s%%" % _fmt_num(ind["multa_percentual_max"], 1)
                       .rstrip("0").rstrip(","), False))
    valores = f.get("valores") or []
    if valores:
        linhas.append(("Valores", "; ".join(
            ("%s: %s" % (v["rotulo"], v["valor"])) if v.get("rotulo") else v["valor"]
            for v in valores[:4]), False))
    if f.get("slas"):
        linhas.append(("SLAs", "; ".join(s["linha"] for s in f["slas"][:3]), False))
    foro = f.get("foro")
    if foro:
        linhas.append(("Foro", foro.get("cidade") or "não preenchido",
                       bool(foro.get("pendente"))))
    for r in f.get("responsaveis", [])[:4]:
        linhas.append((r["papel"], r["nome"], False))
    return linhas


def ocorrencias_visiveis(resultado):
    return [o for o in resultado["ocorrencias"] if not o.get("oculto")]


# ----------------------------------------------------------------------------
# PDF (HTML estatico para xhtml2pdf)
# ----------------------------------------------------------------------------

CSS_PDF = """
@page { size: a4 portrait; margin: 1.4cm 1.3cm 1.6cm 1.3cm;
        @frame rodape { -pdf-frame-content: rodape; bottom: 0.6cm; margin-left: 1.3cm;
                        margin-right: 1.3cm; height: 0.8cm; } }
body { font-family: Helvetica, Arial, sans-serif; font-size: 9.5pt; color: #1a1a1a; }
h1 { font-size: 15pt; margin: 0 0 2px 0; }
h2 { font-size: 11.5pt; margin: 14px 0 6px 0; border-bottom: 1px solid #cccccc;
     padding-bottom: 3px; color: {{FORTE}}; }
.sub { color: #666666; font-size: 8.5pt; margin-bottom: 8px; }
table { width: 100%; margin-bottom: 8px; }
th { background-color: {{CABECALHO}}; color: #ffffff; text-align: left; padding: 4px 6px;
     font-size: 8.5pt; }
td { border-bottom: 1px solid #dddddd; padding: 3px 6px; font-size: 8.5pt;
     vertical-align: top; }
td.n { text-align: center; }
.topo td { border: none; padding: 0 8px 0 0; vertical-align: middle; }
.produto { font-size: 8pt; color: {{FORTE}}; font-weight: bold; }
.meta td { border: none; padding: 2px 6px; }
.meta td.k { color: #666666; width: 120px; }
.alerta-box { background-color: #FBEBD7; color: #8A4B08; padding: 6px 8px;
              font-size: 8.5pt; margin-bottom: 8px; }
.ruim { color: #B00020; font-weight: bold; }
table.oc { margin-bottom: 6px; }
table.oc td.corpo { border: 1px solid #d8d8d8; padding: 5px 8px; }
.cab { font-size: 8pt; color: #555555; margin-bottom: 4px; }
.selo { font-size: 7.5pt; padding: 1px 4px; font-weight: bold; }
.lbl { font-size: 7.5pt; color: #777777; margin: 4px 0 1px 0; }
.txt { margin: 0 0 3px 0; line-height: 1.35; }
.fatores { font-size: 7.5pt; color: #666666; margin-top: 3px; }
.ins { background-color: #d6f5df; color: #0d5c2a; }
.del { background-color: #ffdde1; color: #90001a; text-decoration: line-through; }
.mov { background-color: #e3e3f7; color: #2f2f8f; }
.legenda { font-size: 8pt; color: #555555; margin-bottom: 10px; }
.vazio { color: #888888; font-style: italic; }
#rodape { font-size: 7.5pt; color: #777777; border-top: 1px solid #dddddd; }
"""


def css_pdf(resultado, cfg=None):
    tokens = identidade_mod.gerar_tokens(_ident(resultado, cfg)["cor_primaria"])
    return (CSS_PDF.replace("{{CABECALHO}}", tokens["relatorio_cabecalho"])
            .replace("{{FORTE}}", tokens["claro"]["primaria-forte"]))


def fmt_valor(valor):
    """10.0 -> '10'; 2.5 -> '2,5'; texto fica como esta."""
    if isinstance(valor, float):
        if valor.is_integer():
            return _fmt_num(valor, 0)
        return _fmt_num(valor, 2).rstrip("0").rstrip(",")
    return "-" if valor is None else str(valor)


def _selo_html(texto, cor, fundo="#ffffff"):
    return ("<span class='selo' style='color:%s;background-color:%s;border:1px solid %s'>"
            "%s</span>" % (cor, fundo, cor, html.escape(texto)))


def montar_html_pdf(resultado, cfg=None):
    """HTML estatico (so tabelas e CSS basico) usado para gerar o PDF."""
    e = html.escape
    meta = resultado["meta"]
    ident = _ident(resultado, cfg)
    logo = identidade_mod.logo_data_uri(ident, (cfg or {}).get("caminho_logo") or None)
    p = ["<html><head><meta charset='utf-8'><title>%s</title><style>%s</style></head><body>"
         % (e(ident["nome_produto"]), css_pdf(resultado, cfg))]
    p.append("<div id='rodape'>%s &middot; %s &middot; p&aacute;gina <pdf:pagenumber/> de "
             "<pdf:pagecount/></div>" % (e(texto_rodape(resultado, cfg)),
                                         e(meta.get("fornecedor", ""))))
    p.append("<table class='topo'><tr><td width='46'><img src='%s' width='40' height='40' "
             "alt='Logo %s'/></td><td><div class='produto'>%s</div>"
             "<h1>Comparativo de Devolu&ccedil;&atilde;o Contratual</h1></td></tr></table>"
             % (logo, e(ident["nome_produto"]), e(ident["nome_produto"].upper())))
    p.append("<div class='sub'>Gerado em %s &middot; motor v%s &middot; modelo: %s &middot; "
             "devolu&ccedil;&atilde;o: %s</div>" % (
                 e(meta["data_hora"]), e(meta.get("versao_motor", "")),
                 e(meta["arquivo_modelo"]), e(meta["arquivo_devolucao"])))

    # risco geral + contadores
    cor = CORES_RISCO.get(meta.get("selo_risco"), "#555555")
    p.append("<table class='meta'><tr><td class='k'>Risco geral</td><td>%s &nbsp; "
             "<b style='color:%s'>%d/100</b> &nbsp; &middot; &nbsp; %d cl&aacute;usulas "
             "detectadas &middot; %.1f%% dos par&aacute;grafos id&ecirc;nticos</td></tr></table>"
             % (_selo_html(meta.get("selo_risco", "-"), cor), cor, meta.get("risco_geral", 0),
                meta.get("clausulas_detectadas", 0), meta.get("similaridade_global", 0)))

    p.append("<h2>Ficha do contrato</h2><table class='meta'>")
    for rotulo, valor, ruim in linhas_ficha(resultado):
        p.append("<tr><td class='k'>%s</td><td%s>%s</td></tr>"
                 % (e(rotulo), " class='ruim'" if ruim else "", e(str(valor))))
    p.append("</table>")

    alertas = list((resultado.get("ficha") or {}).get("alertas", []))
    if resultado.get("pendencias"):
        campos = sorted({x["placeholder"] for x in resultado["pendencias"]})
        alertas.insert(0, "Devolução incompleta: %d campo(s) do modelo sem "
                          "preencher — %s" % (len(resultado["pendencias"]),
                                                     ", ".join(campos[:12])))
    for alerta in alertas:
        p.append("<div class='alerta-box'>%s</div>" % e(alerta))

    playbook = resultado.get("playbook") or []
    if playbook:
        p.append("<h2>Playbook de aceita&ccedil;&atilde;o</h2><table><tr>"
                 "<th width='55%'>Regra</th><th width='20%'>Encontrado</th>"
                 "<th width='25%'>Situa&ccedil;&atilde;o</th></tr>")
        for r in playbook:
            situacao = {"ok": ("dentro da política", "#1F7A3D"),
                        "fora": ("FORA DA POLÍTICA", "#B00020"),
                        "sem_dado": ("sem dado", "#777777")}[r["status"]]
            p.append("<tr><td>%s</td><td>%s</td><td style='color:%s'><b>%s</b></td></tr>"
                     % (e(r["rotulo"]), e(fmt_valor(r["encontrado"])),
                        situacao[1], situacao[0]))
        p.append("</table>")

    resumo = resultado["resumo"]
    p.append("<h2>Resumo</h2><table><tr><th width='60%'>Categoria</th>"
             "<th width='20%'>Qtd.</th><th width='20%'>Cr&iacute;ticas</th></tr>")
    ocultas = set(meta.get("categorias_ocultas") or [])
    for tipo in TIPOS:
        if not resumo.get(tipo, {}).get("total") and tipo in ("FORMATACAO",):
            continue
        cor_t, fundo_t = CORES_TIPO[tipo]
        p.append("<tr><td>%s%s</td><td class='n'>%d</td><td class='n'>%d</td></tr>"
                 % (_selo_html(NOMES_TIPO[tipo], cor_t, fundo_t),
                    " <span style='color:#888888'>(oculto no detalhamento)</span>"
                    if tipo in ocultas and resumo[tipo]["total"] else "",
                    resumo[tipo]["total"], resumo[tipo]["criticos"]))
    p.append("<tr><td><b>TOTAL</b></td><td class='n'><b>%d</b></td><td class='n'><b>%d</b>"
             "</td></tr></table>" % (resumo["TOTAL"]["total"], resumo["TOTAL"]["criticos"]))

    por_clausula = resultado.get("por_clausula") or []
    if por_clausula:
        p.append("<h2>Altera&ccedil;&otilde;es por cl&aacute;usula</h2><table><tr>"
                 "<th width='44%'>Cl&aacute;usula</th><th width='8%'>Pre.</th>"
                 "<th width='8%'>Alt.</th><th width='8%'>Inc.</th><th width='8%'>Rem.</th>"
                 "<th width='8%'>Mov.</th><th width='8%'>Pend.</th>"
                 "<th width='8%'>Risco</th></tr>")
        for c in por_clausula:
            p.append("<tr><td>%s</td>%s<td class='n'>%d</td></tr>" % (
                e(c["clausula"]), "".join("<td class='n'>%d</td>" % c.get(t, 0) for t in (
                    "PREENCHIDO", "ALTERADO", "INCLUIDO", "REMOVIDO", "MOVIDO", "PENDENTE")),
                c.get("risco_max", 0)))
        p.append("</table>")

    p.append("<div class='legenda'>Legenda: <span class='del'>texto removido</span> &nbsp; "
             "<span class='ins'>texto inclu&iacute;do</span> &nbsp; "
             "<span class='mov'>bloco movido</span></div>")

    visiveis = ocorrencias_visiveis(resultado)
    p.append("<h2>Detalhamento (%d ocorr&ecirc;ncias)</h2>" % len(visiveis))
    if not visiveis:
        p.append("<p class='vazio'>Nenhuma diferen&ccedil;a de conte&uacute;do identificada "
                 "entre os documentos.</p>")
    for oc in visiveis:
        p.append(_item_pdf(oc, e))
    p.append("</body></html>")
    return "\n".join(p)


def _item_pdf(oc, e):
    cor_t, fundo_t = CORES_TIPO.get(oc["tipo"], ("#555555", "#eeeeee"))
    cor_r = CORES_RISCO.get(oc["selo"], "#555555")
    selos = [_selo_html(NOMES_TIPO[oc["tipo"]].upper(), cor_t, fundo_t),
             _selo_html("Risco %s %d" % (oc["selo"], oc["risco"]), cor_r)]
    if oc.get("critico"):
        selos.append(_selo_html("CRÍTICO", "#7a5b00", "#ffe9a8"))
    if oc.get("tambem_movido"):
        selos.append(_selo_html("movido", "#4A4AB5", "#E3E3F7"))
    for regra in oc.get("fora_politica", []):
        selos.append(_selo_html("fora da política: " + regra, "#B00020", "#FFDDE1"))
    if oc.get("aprovado_antes"):
        ap = oc["aprovado_antes"]
        selos.append(_selo_html("já aprovado em %s%s" % (
            ap.get("data", ""), (" / " + ap["fornecedor"]) if ap.get("fornecedor") else ""),
            "#1F7A3D", "#D6F5DF"))
    p = ["<table class='oc'><tr><td class='corpo' style='border-left: 4px solid %s'>" % cor_t,
         "<div class='cab'>#%s &middot; %s &middot; %s &middot; modelo: %s | "
         "devolu&ccedil;&atilde;o: %s</div>" % (
             oc["id"].split("-")[-1], " ".join(selos), e(oc["clausula"]),
             oc["pos_modelo"] or "-", oc["pos_devolucao"] or "-")]
    if oc["tipo"] in ("PREENCHIDO", "PENDENTE"):
        if oc.get("campos"):
            p.append("<div class='lbl'>CAMPOS PREENCHIDOS</div><p class='txt'>%s</p>" %
                     "<br/>".join("%s &rarr; <b>%s</b>" % (e(c["placeholder"]), e(c["valor"]))
                                  for c in oc["campos"]))
        if oc.get("pendentes"):
            p.append("<div class='lbl'>CAMPOS SEM PREENCHER</div><p class='txt' "
                     "style='color:#8A4B08'>%s</p>" % e(", ".join(oc["pendentes"])))
        p.append("<div class='lbl'>TEXTO NA DEVOLU&Ccedil;&Atilde;O</div><p class='txt'>%s</p>"
                 % e(oc["texto_devolucao"]))
    else:
        rot_a, rot_b = ROTULOS_BLOCO.get(oc["tipo"], ("MODELO", "DEVOLUÇÃO"))
        if rot_a:
            p.append("<div class='lbl'>%s</div><p class='txt'>%s</p>"
                     % (e(rot_a), oc["html_modelo"]))
        if rot_b:
            p.append("<div class='lbl'>%s</div><p class='txt'>%s</p>"
                     % (e(rot_b), oc["html_devolucao"]))
    if oc.get("fatores") and oc["tipo"] not in ("PREENCHIDO", "FORMATACAO"):
        p.append("<div class='fatores'>Risco: %s</div>" % e("; ".join(oc["fatores"])))
    p.append("</td></tr></table>")
    return "".join(p)


def gravar_pdf(html_str, caminho, log):
    try:
        logging.getLogger("xhtml2pdf").setLevel(logging.ERROR)
        from xhtml2pdf import pisa
    except ImportError:
        log.warning("xhtml2pdf nao instalado - PDF nao gerado (pip install xhtml2pdf)")
        return False
    from io import BytesIO
    origem = BytesIO(html_str.encode("utf-8"))
    tmp = caminho + ".tmp"
    with open(tmp, "wb") as destino:
        status = pisa.CreatePDF(origem, dest=destino, encoding="utf-8")
    if status.err:
        log.warning("Falha ao gerar PDF: %s", caminho)
        os.remove(tmp)
        return False
    os.replace(tmp, caminho)
    return True


# ----------------------------------------------------------------------------
# HTML interativo
# ----------------------------------------------------------------------------

def montar_html_interativo(resultado, cfg=None):
    import comparativo_interativo
    return comparativo_interativo.montar(resultado, cfg)


# ----------------------------------------------------------------------------
# TXT
# ----------------------------------------------------------------------------

def montar_txt(resultado, cfg=None):
    meta, resumo = resultado["meta"], resultado["resumo"]
    ident = _ident(resultado, cfg)
    sep = "=" * 78
    linhas = [sep, ident["nome_produto"].upper(), "COMPARATIVO DE DEVOLUCAO CONTRATUAL", sep,
              "Gerado em.........: %s" % meta["data_hora"],
              "Modelo padrao.....: %s" % meta["arquivo_modelo"],
              "Devolucao.........: %s" % meta["arquivo_devolucao"],
              "Paragrafos identicos: %.1f%%" % meta["similaridade_global"],
              "Risco geral.......: %d/100 (%s)" % (meta.get("risco_geral", 0),
                                                   meta.get("selo_risco", "-")),
              "", "FICHA DO CONTRATO", "-" * 78]
    for rotulo, valor, ruim in linhas_ficha(resultado):
        linhas.append("  %-18s %s%s" % (rotulo[:18] + ":", valor, "  <<< VERIFICAR" if ruim
                                        else ""))
    alertas = list((resultado.get("ficha") or {}).get("alertas", []))
    if resultado.get("pendencias"):
        alertas.insert(0, "DEVOLUCAO INCOMPLETA - campos sem preencher: %s" % ", ".join(
            sorted({p["placeholder"] for p in resultado["pendencias"]})))
    if alertas:
        linhas += ["", "ALERTAS", "-" * 78] + ["  ! %s" % a for a in alertas]
    if resultado.get("playbook"):
        linhas += ["", "PLAYBOOK DE ACEITACAO", "-" * 78]
        for r in resultado["playbook"]:
            linhas.append("  [%s] %s (encontrado: %s)" % (
                {"ok": " OK ", "fora": "FORA", "sem_dado": " -- "}[r["status"]], r["rotulo"],
                fmt_valor(r["encontrado"])))
    linhas += ["", "RESUMO", "-" * 78]
    for tipo in TIPOS:
        linhas.append("  %-11s %4d  (criticos: %d)" % (tipo, resumo[tipo]["total"],
                                                       resumo[tipo]["criticos"]))
    linhas += ["  %-11s %4d  (criticos: %d)" % ("TOTAL", resumo["TOTAL"]["total"],
                                                resumo["TOTAL"]["criticos"]), ""]
    if resultado.get("por_clausula"):
        linhas += ["POR CLAUSULA", "-" * 78]
        for c in resultado["por_clausula"]:
            linhas.append("  %-52.52s %s  risco max %d" % (c["clausula"], " ".join(
                "%s%d" % (t[:3], c.get(t, 0)) for t in ("PREENCHIDO", "ALTERADO", "INCLUIDO",
                                                        "REMOVIDO", "PENDENTE")
                if c.get(t)), c.get("risco_max", 0)))
        linhas.append("")
    linhas += ["Legenda: [--texto suprimido--]   {++texto incluido++}", "",
               "DETALHAMENTO", "-" * 78]
    for oc in ocorrencias_visiveis(resultado):
        flag = "  <<< CRITICO" if oc.get("critico") else ""
        linhas += ["", "#%s  [%s]  %s  | risco %d (%s)%s" % (
            oc["id"].split("-")[-1], oc["tipo"], oc["clausula"], oc["risco"], oc["selo"], flag),
            "    modelo: %s | devolucao: %s | similaridade: %s%%"
            % (oc["pos_modelo"] or "-", oc["pos_devolucao"] or "-", oc["similaridade"])]
        if oc["tipo"] in ("PREENCHIDO", "PENDENTE"):
            for c in oc.get("campos", []):
                linhas.append("    CAMPO    : %s -> %s" % (c["placeholder"], c["valor"]))
            if oc.get("pendentes"):
                linhas.append("    PENDENTE : %s" % ", ".join(oc["pendentes"]))
        elif oc["tipo"] in ("ALTERADO", "FORMATACAO"):
            linhas += ["    ORIGINAL : %s" % oc["texto_modelo"],
                       "    DEVOLVIDO: %s" % oc["texto_devolucao"],
                       "    DIFF     : %s" % oc["marcado"]]
        elif oc["tipo"] == "REMOVIDO":
            linhas.append("    SUPRIMIDO: %s" % oc["texto_modelo"])
        else:
            linhas.append("    TEXTO    : %s" % oc["texto_devolucao"])
        if oc.get("fatores") and oc["tipo"] not in ("PREENCHIDO", "FORMATACAO"):
            linhas.append("    RISCO    : %s" % "; ".join(oc["fatores"]))
    linhas += ["", sep, texto_rodape(resultado, cfg)]
    return "\n".join(linhas)


# ----------------------------------------------------------------------------
# XLSX
# ----------------------------------------------------------------------------

def gravar_xlsx(caminho, resultado, cfg, log):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.warning("openpyxl nao instalado - XLSX nao gerado")
        return False

    meta, resumo = resultado["meta"], resultado["resumo"]
    ident = _ident(resultado, cfg)
    tokens = identidade_mod.gerar_tokens(ident["cor_primaria"])
    cab_fill = PatternFill("solid", fgColor=tokens["relatorio_cabecalho"].lstrip("#"))
    cab_font = Font(color="FFFFFF", bold=True)
    quebra = Alignment(wrap_text=True, vertical="top")

    def cabecalho(ws, linha):
        for cel in ws[linha]:
            if cel.value is not None:
                cel.fill, cel.font = cab_fill, cab_font

    def larguras(ws, valores):
        for i, largura in enumerate(valores, 1):
            ws.column_dimensions[get_column_letter(i)].width = largura

    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMO"
    ws.append([ident["nome_produto"], "motor v%s" % meta.get("versao_motor", "")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Campo", "Valor"])
    cabecalho(ws, 2)
    for chave, valor in [
        ("Gerado em", meta["data_hora"]), ("Fornecedor", meta["fornecedor"]),
        ("CNPJ", meta["cnpj"]), ("Modelo padrao", meta["arquivo_modelo"]),
        ("Devolucao", meta["arquivo_devolucao"]),
        ("Paragrafos modelo", meta["qtd_modelo"]),
        ("Paragrafos devolucao", meta["qtd_devolucao"]),
        ("Paragrafos identicos (%)", round(meta["similaridade_global"], 1)),
        ("Risco geral (0-100)", meta.get("risco_geral")),
        ("Nivel de risco", meta.get("selo_risco")),
        ("Clausulas detectadas", meta.get("clausulas_detectadas")),
        ("Devolucao incompleta", "SIM" if meta.get("devolucao_incompleta") else "nao"),
        ("Regras fora da politica", meta.get("fora_da_politica", 0)),
    ]:
        ws.append([chave, valor])
    ws.append([])
    ws.append(["Categoria", "Quantidade", "Criticas"])
    cabecalho(ws, ws.max_row)
    for tipo in TIPOS + ("TOTAL",):
        ws.append([tipo, resumo[tipo]["total"], resumo[tipo]["criticos"]])
        if tipo in CORES_TIPO:
            cor, fundo = CORES_TIPO[tipo]
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=fundo.lstrip("#"))
            ws.cell(ws.max_row, 1).font = Font(color=cor.lstrip("#"), bold=True)
    if resultado.get("por_clausula"):
        ws.append([])
        ws.append(["Clausula"] + list(TIPOS[:6]) + ["Risco maximo"])
        cabecalho(ws, ws.max_row)
        for c in resultado["por_clausula"]:
            ws.append([c["clausula"]] + [c.get(t, 0) for t in TIPOS[:6]]
                      + [c.get("risco_max", 0)])
    larguras(ws, [46, 22, 14, 14, 14, 14, 14, 14])

    wf = wb.create_sheet("FICHA")
    wf.append(["Campo", "Valor", "Verificar"])
    cabecalho(wf, 1)
    for rotulo, valor, ruim in linhas_ficha(resultado):
        wf.append([rotulo, valor, "SIM" if ruim else ""])
    ficha = resultado.get("ficha") or {}
    for titulo, itens, colunas in (
            ("CNPJs encontrados", ficha.get("cnpjs", []), ("numero", "valido", "contexto")),
            ("Valores", ficha.get("valores", []), ("valor", "rotulo", "contexto")),
            ("Percentuais", ficha.get("percentuais", []), ("valor", "contexto")),
            ("Prazos", ficha.get("prazos", []), ("valor", "horas", "contexto")),
            ("SLAs", ficha.get("slas", []), ("indicador", "horas", "linha"))):
        if not itens:
            continue
        wf.append([])
        wf.append([titulo] + [c.capitalize() for c in colunas])
        cabecalho(wf, wf.max_row)
        for item in itens:
            wf.append([""] + [("SIM" if item.get(c) is True else "NAO"
                               if item.get(c) is False else item.get(c)) for c in colunas])
    if ficha.get("alertas"):
        wf.append([])
        wf.append(["Alertas"])
        cabecalho(wf, wf.max_row)
        for a in ficha["alertas"]:
            wf.append([a])
    larguras(wf, [28, 60, 70, 60])
    for linha in wf.iter_rows(min_row=2):
        for cel in linha:
            cel.alignment = quebra

    wd = wb.create_sheet("DETALHE")
    wd.append(["#", "Id", "Tipo", "Risco", "Nivel", "Critico", "Clausula", "Pos. modelo",
               "Pos. devolucao", "Similaridade (%)", "Campos preenchidos", "Pendentes",
               "Texto do modelo", "Texto da devolucao", "Diff marcado", "Fatores de risco",
               "Oculto por padrao"])
    cabecalho(wd, 1)
    for n, a in enumerate(resultado["ocorrencias"], 1):
        wd.append([n, a["id"], a["tipo"], a["risco"], a["selo"],
                   "SIM" if a.get("critico") else "", a["clausula"], a["pos_modelo"],
                   a["pos_devolucao"], a["similaridade"],
                   "; ".join("%s = %s" % (c["placeholder"], c["valor"])
                             for c in a.get("campos", [])),
                   ", ".join(a.get("pendentes", [])), a["texto_modelo"],
                   a["texto_devolucao"], a["marcado"], "; ".join(a.get("fatores", [])),
                   "SIM" if a.get("oculto") else ""])
        cor, fundo = CORES_TIPO.get(a["tipo"], ("000000", "FFFFFF"))
        wd.cell(wd.max_row, 3).fill = PatternFill("solid", fgColor=fundo.lstrip("#"))
        wd.cell(wd.max_row, 3).font = Font(color=cor.lstrip("#"), bold=True)
    larguras(wd, [5, 8, 13, 7, 8, 8, 30, 10, 12, 12, 40, 26, 60, 60, 60, 50, 10])
    for linha in wd.iter_rows(min_row=2, min_col=11, max_col=16):
        for cel in linha:
            cel.alignment = quebra
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = wd.dimensions

    if resultado.get("pendencias"):
        wp = wb.create_sheet("PENDENCIAS")
        wp.append(["Campo sem preencher", "Clausula", "Ocorrencia", "Pos. devolucao"])
        cabecalho(wp, 1)
        for x in resultado["pendencias"]:
            wp.append([x["placeholder"], x["clausula"], x["ocorrencia"], x["pos_devolucao"]])
        larguras(wp, [34, 46, 12, 14])

    if resultado.get("playbook"):
        wpl = wb.create_sheet("PLAYBOOK")
        wpl.append(["Regra", "Campo", "Operador", "Esperado", "Encontrado", "Situacao"])
        cabecalho(wpl, 1)
        for r in resultado["playbook"]:
            wpl.append([r["rotulo"], r["campo"], r["operador"], r["esperado"],
                        r["encontrado"], {"ok": "dentro", "fora": "FORA",
                                          "sem_dado": "sem dado"}[r["status"]]])
        larguras(wpl, [40, 24, 10, 14, 14, 12])

    tmp = caminho + ".tmp"
    wb.save(tmp)
    os.replace(tmp, caminho)
    return True
