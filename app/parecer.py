# -*- coding: utf-8 -*-
"""
parecer.py
Parecer de revisao (PDF e DOCX) para aprovacao interna: ficha do contrato, pontos
fora da politica, pendencias, decisoes da revisao com comentarios e conclusao.
"""

import os
import html
from datetime import datetime

import identidade as identidade_mod
import relatorios

STATUS_REVISAO = {"pendente": "Pendente", "em_revisao": "Em revisão",
                  "aprovado": "Aprovado", "devolvido": "Devolvido ao fornecedor"}
CONCLUSOES = {
    "aprovado": "Revisão concluída com recomendação de APROVAÇÃO "
                "da devolução, nos termos das decisões abaixo.",
    "devolvido": "Revisão concluída com DEVOLUÇÃO ao fornecedor: os pontos "
                 "rejeitados seguem na contraproposta com controle de alterações.",
    "em_revisao": "Revisão em andamento. Este parecer é parcial.",
    "pendente": "Revisão ainda não iniciada. Este parecer traz apenas a análise "
                "automática.",
}


def montar_dados(resultado, revisao, execucao, observacao=""):
    itens = (revisao or {}).get("itens", {})
    decisoes, pendentes_criticas = [], []
    for oc in resultado.get("ocorrencias", []):
        it = itens.get(oc.get("chave"), {})
        if it.get("status") or it.get("comentario"):
            decisoes.append((oc, it))
        elif oc.get("critico"):
            pendentes_criticas.append(oc)
    decisoes.sort(key=lambda x: (-x[0].get("risco", 0)))
    status = (execucao or {}).get("status_revisao", "pendente")
    return {
        "meta": resultado["meta"],
        "ficha": relatorios.linhas_ficha(resultado),
        "fora": [r for r in resultado.get("playbook") or [] if r["status"] == "fora"],
        "pendencias": sorted({p["placeholder"] for p in resultado.get("pendencias", [])}),
        "alertas": (resultado.get("ficha") or {}).get("alertas", []),
        "decisoes": decisoes,
        "pendentes_criticas": pendentes_criticas,
        "aceitas": sum(1 for _, it in decisoes if it.get("status") == "aceito"),
        "rejeitadas": sum(1 for _, it in decisoes if it.get("status") == "rejeitado"),
        "status": status,
        "status_rotulo": STATUS_REVISAO.get(status, status),
        "responsavel": (execucao or {}).get("responsavel") or "—",
        "departamento": (execucao or {}).get("nome_departamento") or "",
        "conclusao": CONCLUSOES.get(status, ""),
        "observacao": (observacao or "").strip()[:3000],
        "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
    }


def _decisao(it):
    return {"aceito": "Aceita", "rejeitado": "Rejeitada"}.get(it.get("status"), "Comentada")


def html_pdf(resultado, revisao, execucao, cfg=None, observacao=""):
    e = html.escape
    d = montar_dados(resultado, revisao, execucao, observacao)
    ident = identidade_mod.mesclar((cfg or {}).get("identidade") or resultado.get("identidade"))
    logo = identidade_mod.logo_data_uri(ident, (cfg or {}).get("caminho_logo") or None)
    m = d["meta"]
    p = ["<html><head><meta charset='utf-8'><style>%s</style></head><body>"
         % relatorios.css_pdf(resultado, cfg)]
    p.append("<div id='rodape'>%s &middot; Parecer gerado em %s &middot; p&aacute;gina "
             "<pdf:pagenumber/> de <pdf:pagecount/></div>"
             % (e(relatorios.texto_rodape(resultado, cfg)), d["gerado_em"]))
    p.append("<table class='topo'><tr><td width='46'><img src='%s' width='40' height='40'/>"
             "</td><td><div class='produto'>%s</div><h1>Parecer de revis&atilde;o "
             "contratual</h1></td></tr></table>" % (logo, e(ident["nome_produto"].upper())))
    p.append("<table class='meta'>")
    for k, v in (("Fornecedor", m.get("fornecedor")), ("CNPJ", m.get("cnpj") or "—"),
                 ("Departamento", d["departamento"]), ("Modelo", m.get("arquivo_modelo")),
                 ("Devolução", m.get("arquivo_devolucao")),
                 ("Análise", "%s · motor v%s" % (m.get("data_hora"),
                                                              m.get("versao_motor"))),
                 ("Situação", d["status_rotulo"]), ("Responsável", d["responsavel"]),
                 ("Risco geral", "%s/100 (%s)" % (m.get("risco_geral"), m.get("selo_risco")))):
        p.append("<tr><td class='k'>%s</td><td>%s</td></tr>" % (e(k), e(str(v or "—"))))
    p.append("</table>")
    p.append("<h2>Conclus&atilde;o</h2><p>%s</p>" % e(d["conclusao"]))
    if d["observacao"]:
        p.append("<p><b>Observa&ccedil;&otilde;es do revisor:</b> %s</p>"
                 % e(d["observacao"]).replace("\n", "<br/>"))
    p.append("<p class='sub'>%d decis&atilde;o(&otilde;es): %d aceita(s), %d rejeitada(s). "
             "%d cr&iacute;tica(s) ainda sem decis&atilde;o.</p>"
             % (len(d["decisoes"]), d["aceitas"], d["rejeitadas"], len(d["pendentes_criticas"])))

    p.append("<h2>Ficha do contrato</h2><table class='meta'>")
    for rotulo, valor, ruim in d["ficha"]:
        p.append("<tr><td class='k'>%s</td><td%s>%s</td></tr>"
                 % (e(rotulo), " class='ruim'" if ruim else "", e(str(valor))))
    p.append("</table>")
    if d["fora"] or d["pendencias"] or d["alertas"]:
        p.append("<h2>Pontos de aten&ccedil;&atilde;o</h2>")
        for r in d["fora"]:
            p.append("<div class='alerta-box'>Fora da pol&iacute;tica: %s (encontrado: %s)"
                     "</div>" % (e(r["rotulo"]), e(relatorios.fmt_valor(r["encontrado"]))))
        if d["pendencias"]:
            p.append("<div class='alerta-box'>Devolu&ccedil;&atilde;o incompleta: %s</div>"
                     % e(", ".join(d["pendencias"])))
        for a in d["alertas"]:
            p.append("<div class='alerta-box'>%s</div>" % e(a))
    if d["decisoes"]:
        p.append("<h2>Decis&otilde;es da revis&atilde;o</h2><table><tr><th width='6%'>#</th>"
                 "<th width='14%'>Tipo</th><th width='28%'>Cl&aacute;usula</th>"
                 "<th width='14%'>Decis&atilde;o</th><th width='38%'>Coment&aacute;rio</th></tr>")
        for oc, it in d["decisoes"]:
            p.append("<tr><td>%s</td><td>%s</td><td>%s</td><td><b>%s</b></td><td>%s</td></tr>"
                     % (oc["id"].split("-")[-1], e(relatorios.NOMES_TIPO[oc["tipo"]]),
                        e(oc["clausula"]), _decisao(it), e(it.get("comentario", ""))))
        p.append("</table>")
    if d["pendentes_criticas"]:
        p.append("<h2>Cr&iacute;ticas sem decis&atilde;o</h2><table><tr><th width='6%'>#</th>"
                 "<th width='14%'>Tipo</th><th width='30%'>Cl&aacute;usula</th>"
                 "<th width='50%'>Trecho</th></tr>")
        for oc in d["pendentes_criticas"]:
            trecho = (oc["texto_devolucao"] or oc["texto_modelo"])[:220]
            p.append("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>"
                     % (oc["id"].split("-")[-1], e(relatorios.NOMES_TIPO[oc["tipo"]]),
                        e(oc["clausula"]), e(trecho)))
        p.append("</table>")
    p.append("<br/><br/><table class='meta'><tr><td width='50%'>_______________________________"
             "<br/>Respons&aacute;vel pela revis&atilde;o</td><td>_______________________________"
             "<br/>Aprova&ccedil;&atilde;o</td></tr></table>")
    p.append("</body></html>")
    return "\n".join(p)


def gravar_docx(caminho, resultado, revisao, execucao, cfg=None, observacao=""):
    from docx import Document
    from docx.shared import Pt, RGBColor
    d = montar_dados(resultado, revisao, execucao, observacao)
    ident = identidade_mod.mesclar((cfg or {}).get("identidade") or resultado.get("identidade"))
    tokens = identidade_mod.gerar_tokens(ident["cor_primaria"])
    cor = RGBColor.from_string(tokens["relatorio_cabecalho"].lstrip("#"))
    m = d["meta"]
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    rotulo = doc.add_paragraph()
    run = rotulo.add_run(ident["nome_produto"].upper())
    run.bold, run.font.size, run.font.color.rgb = True, Pt(8), cor
    doc.add_heading("Parecer de revisão contratual", level=1)

    tabela = doc.add_table(rows=0, cols=2)
    for k, v in (("Fornecedor", m.get("fornecedor")), ("CNPJ", m.get("cnpj") or "—"),
                 ("Departamento", d["departamento"]), ("Modelo", m.get("arquivo_modelo")),
                 ("Devolução", m.get("arquivo_devolucao")),
                 ("Situação", d["status_rotulo"]),
                 ("Responsável", d["responsavel"]),
                 ("Risco geral", "%s/100 (%s)" % (m.get("risco_geral"), m.get("selo_risco")))):
        c = tabela.add_row().cells
        c[0].text, c[1].text = k, str(v or "—")
    doc.add_heading("Conclusão", level=2)
    doc.add_paragraph(d["conclusao"])
    if d["observacao"]:
        doc.add_paragraph("Observações do revisor: " + d["observacao"])
    doc.add_paragraph("%d decisão(ões): %d aceita(s), %d rejeitada(s); %d "
                      "crítica(s) sem decisão." % (len(d["decisoes"]), d["aceitas"],
                                                              d["rejeitadas"],
                                                              len(d["pendentes_criticas"])))
    doc.add_heading("Ficha do contrato", level=2)
    tabela = doc.add_table(rows=0, cols=2)
    for rotulo_f, valor, _ in d["ficha"]:
        c = tabela.add_row().cells
        c[0].text, c[1].text = rotulo_f, str(valor)
    if d["fora"] or d["pendencias"] or d["alertas"]:
        doc.add_heading("Pontos de atenção", level=2)
        for r in d["fora"]:
            doc.add_paragraph("Fora da política: %s (encontrado: %s)"
                              % (r["rotulo"], relatorios.fmt_valor(r["encontrado"])),
                              style="List Bullet")
        if d["pendencias"]:
            doc.add_paragraph("Devolução incompleta: " + ", ".join(d["pendencias"]),
                              style="List Bullet")
        for a in d["alertas"]:
            doc.add_paragraph(a, style="List Bullet")
    if d["decisoes"]:
        doc.add_heading("Decisões da revisão", level=2)
        tabela = doc.add_table(rows=1, cols=4)
        tabela.style = "Table Grid"
        for c, t in zip(tabela.rows[0].cells, ("#", "Cláusula", "Decisão",
                                               "Comentário")):
            c.text = t
        for oc, it in d["decisoes"]:
            c = tabela.add_row().cells
            c[0].text = oc["id"].split("-")[-1]
            c[1].text = "%s — %s" % (relatorios.NOMES_TIPO[oc["tipo"]], oc["clausula"])
            c[2].text = _decisao(it)
            c[3].text = it.get("comentario", "")
    if d["pendentes_criticas"]:
        doc.add_heading("Críticas sem decisão", level=2)
        for oc in d["pendentes_criticas"]:
            doc.add_paragraph("#%s %s — %s: %s" % (
                oc["id"].split("-")[-1], relatorios.NOMES_TIPO[oc["tipo"]], oc["clausula"],
                (oc["texto_devolucao"] or oc["texto_modelo"])[:220]), style="List Bullet")
    rodape = doc.sections[0].footer.paragraphs[0]
    rodape.text = "%s · parecer gerado em %s" % (relatorios.texto_rodape(resultado, cfg),
                                                      d["gerado_em"])
    tmp = caminho + ".tmp"
    doc.save(tmp)
    os.replace(tmp, caminho)
    return caminho


def gerar(pasta, resultado, revisao, execucao, cfg=None, observacao="", log=None):
    """Grava parecer.pdf e parecer.docx na pasta do resultado."""
    import logging
    log = log or logging.getLogger("parecer")
    arquivos = []
    if relatorios.gravar_pdf(html_pdf(resultado, revisao, execucao, cfg, observacao),
                             os.path.join(pasta, "parecer.pdf"), log):
        arquivos.append("parecer.pdf")
    gravar_docx(os.path.join(pasta, "parecer.docx"), resultado, revisao, execucao, cfg,
                observacao)
    arquivos.append("parecer.docx")
    return arquivos
