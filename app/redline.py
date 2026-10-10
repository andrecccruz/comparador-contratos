# -*- coding: utf-8 -*-
"""
redline.py
Gera a contraproposta (.docx) com CONTROLE DE ALTERACOES do Word (w:ins / w:del)
a partir da devolucao e das decisoes da revisao:

  - ocorrencia REJEITADA  -> volta ao texto do modelo (alteracao rastreada)
      ALTERADO/FORMATACAO/PREENCHIDO/PENDENTE: paragrafo/celulas reescritos
      INCLUIDO: paragrafo/linha marcado como excluido
      REMOVIDO: paragrafo/linha do modelo reinserido
      MOVIDO:   excluido na posicao nova e reinserido na posicao do modelo
  - ocorrencia ACEITA ou sem decisao -> texto da devolucao mantido
  - comentarios da revisao viram comentarios do Word

A formatacao interna do paragrafo reescrito segue a do primeiro trecho (run).
Devolucoes em PDF/TXT geram um .docx novo com o texto da devolucao.
"""

import os
import copy
from datetime import datetime

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

import comparar_contratos as motor

NOME_ARQUIVO = "contraproposta_redline.docx"


class ErroRedline(Exception):
    pass


class Marcador:
    """Cria elementos w:ins / w:del com id, autor e data."""

    def __init__(self, autor):
        self.autor = autor[:60] or "Revisão"
        self.data = datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")
        self._id = 9000

    def _proximo(self):
        self._id += 1
        return str(self._id)

    def _marca(self, tag):
        el = OxmlElement(tag)
        el.set(qn("w:id"), self._proximo())
        el.set(qn("w:author"), self.autor)
        el.set(qn("w:date"), self.data)
        return el

    def run(self, texto, rpr=None, excluido=False):
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = OxmlElement("w:delText" if excluido else "w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = texto
        r.append(t)
        return r

    def ins(self, texto, rpr=None):
        el = self._marca("w:ins")
        el.append(self.run(texto, rpr))
        return el

    def dele(self, texto, rpr=None):
        el = self._marca("w:del")
        el.append(self.run(texto, rpr, excluido=True))
        return el

    def marcar_paragrafo(self, p, tipo):
        """Marca a quebra do paragrafo como inserida/excluida (w:pPr/w:rPr/w:ins|del)."""
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p.insert(0, ppr)
        rpr = ppr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ppr.append(rpr)
        rpr.append(self._marca("w:ins" if tipo == "ins" else "w:del"))

    def marcar_linha(self, tr, tipo):
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            tr.insert(0, trpr)
        trpr.append(self._marca("w:ins" if tipo == "ins" else "w:del"))


# ----------------------------------------------------------------- operacoes

def _primeiro_rpr(p):
    for r in p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            return rpr
    return None


def _limpar_conteudo(p):
    for filho in list(p):
        if filho.tag != qn("w:pPr"):
            p.remove(filho)


def reescrever_paragrafo(p, texto_atual, texto_alvo, m):
    """Troca o conteudo do paragrafo pelo diff rastreado atual -> alvo."""
    rpr = _primeiro_rpr(p)
    _limpar_conteudo(p)
    for op, trecho in motor.diff_ops(texto_atual, texto_alvo):
        if not trecho:
            continue
        if op == "=":
            p.append(m.run(trecho, rpr))
        elif op == "-":
            p.append(m.dele(trecho, rpr))
        else:
            p.append(m.ins(trecho, rpr))


def excluir_paragrafo(p, m):
    """Todo o texto do paragrafo vira exclusao rastreada."""
    texto = "".join(t.text or "" for t in p.iter(qn("w:t")))
    rpr = _primeiro_rpr(p)
    _limpar_conteudo(p)
    if texto:
        p.append(m.dele(texto, rpr))
    m.marcar_paragrafo(p, "del")


def novo_paragrafo_apos(ancora, texto, m, modelo_ppr=None):
    p = OxmlElement("w:p")
    if modelo_ppr is not None:
        p.append(copy.deepcopy(modelo_ppr))
    p.append(m.ins(texto, _primeiro_rpr(ancora) if ancora is not None else None))
    m.marcar_paragrafo(p, "ins")
    if ancora is not None:
        ancora.addnext(p)
    return p


def _paragrafos_da_celula(tc):
    return tc.findall(qn("w:p"))


def reescrever_celulas(tcs, atuais, alvos, m):
    for tc, atual, alvo in zip(tcs, atuais, alvos):
        if motor.normalizar(atual) == motor.normalizar(alvo):
            continue
        paragrafos = _paragrafos_da_celula(tc)
        if not paragrafos:
            continue
        texto_atual = "\n".join("".join(t.text or "" for t in p.iter(qn("w:t")))
                                for p in paragrafos)
        reescrever_paragrafo(paragrafos[0], " ".join(texto_atual.split("\n")).strip(), alvo, m)
        for extra in paragrafos[1:]:
            excluir_paragrafo(extra, m)


def excluir_linha(tr, m):
    for p in tr.iter(qn("w:p")):
        excluir_paragrafo(p, m)
    m.marcar_linha(tr, "del")


def nova_linha_apos(tr_ancora, celulas, m):
    tr = copy.deepcopy(tr_ancora)
    tcs = tr.findall(qn("w:tc"))
    for tc, texto in zip(tcs, celulas + [""] * (len(tcs) - len(celulas))):
        paragrafos = _paragrafos_da_celula(tc)
        for extra in paragrafos[1:]:
            tc.remove(extra)
        p = paragrafos[0] if paragrafos else OxmlElement("w:p")
        rpr = _primeiro_rpr(p)
        _limpar_conteudo(p)
        if texto:
            p.append(m.ins(texto, rpr))
        m.marcar_paragrafo(p, "ins")
        if not paragrafos:
            tc.append(p)
    m.marcar_linha(tr, "ins")
    tr_ancora.addnext(tr)
    return tr


# ----------------------------------------------------------------- montagem

def _cfg_do_resultado(resultado):
    cfg = dict(motor.CONFIG_PADRAO)
    cfg.update(resultado.get("config_relevante") or {})
    return cfg


def _documento_devolucao(caminho, cfg):
    """(Document, blocos com elementos) — alinhados com as posicoes do resultado."""
    if caminho.lower().endswith(".docx"):
        doc = Document(caminho)
        estrutura = motor.ler_docx_estruturado(caminho, com_elementos=True, documento=doc)
    else:
        doc = Document()
        estrutura = []
        for bloco in motor.ler_documento_estruturado(caminho):
            par = doc.add_paragraph(bloco["texto"])
            estrutura.append(dict(bloco, _el=par._p))
    return doc, motor.montar_blocos(estrutura, cfg)


def _celulas_do_modelo(caminho_modelo, cfg):
    if not caminho_modelo or not caminho_modelo.lower().endswith(".docx"):
        return {}
    blocos = motor.montar_blocos(motor.ler_docx_estruturado(caminho_modelo), cfg)
    return {b["pos"]: b.get("celulas") for b in blocos if b.get("tipo") == "linha_tabela"}


def _ancora_para_modelo(resultado, pos_modelo):
    """Ultima posicao da devolucao alinhada antes da posicao do modelo."""
    melhor = None
    for pm, pd in resultado.get("alinhamento") or []:
        if pm is not None and pm >= pos_modelo:
            break
        if pd is not None:
            melhor = pd
    return melhor


def _elemento_bloco(bloco):
    if bloco.get("_el") is not None:
        return bloco["_el"]
    if bloco.get("_celulas"):
        ps = _paragrafos_da_celula(bloco["_celulas"][0])
        return ps[0] if ps else None
    return None


def gerar(resultado, revisao, caminho_devolucao, caminho_destino, caminho_modelo=None,
          autor="Revisão"):
    """Gera o .docx de contraproposta. Retorna um resumo do que foi aplicado."""
    cfg = _cfg_do_resultado(resultado)
    doc, blocos = _documento_devolucao(caminho_devolucao, cfg)
    por_pos = {b["pos"]: b for b in blocos}
    esperado = {d["pos"]: d["texto"] for d in (resultado.get("documento") or {}).get(
        "devolucao", [])}
    divergentes = [p for p, t in esperado.items()
                   if p not in por_pos or por_pos[p]["texto"] != t]
    if esperado and divergentes:
        raise ErroRedline("O arquivo da devolução não corresponde ao resultado "
                          "(posições %s). Reprocesse a comparação."
                          % ", ".join(map(str, divergentes[:5])))
    celulas_modelo = _celulas_do_modelo(caminho_modelo, cfg)
    itens = (revisao or {}).get("itens", {})
    m = Marcador(autor)
    resumo = {"revertidas": 0, "reinseridas": 0, "excluidas": 0, "mantidas": 0,
              "comentarios": 0, "ignoradas": []}
    comentarios = []

    # insercoes primeiro guardam a ancora original; processa do fim para o comeco
    ocorrencias = sorted(resultado.get("ocorrencias", []),
                         key=lambda o: -(o["pos_devolucao"] or 0))
    for oc in ocorrencias:
        decisao = itens.get(oc.get("chave"), {})
        if decisao.get("comentario"):
            comentarios.append((oc, decisao))
        if decisao.get("status") != "rejeitado":
            resumo["mantidas"] += 1
            continue
        tipo = oc["tipo"]
        bloco = por_pos.get(oc["pos_devolucao"]) if oc["pos_devolucao"] != "" else None
        try:
            if tipo in ("ALTERADO", "FORMATACAO", "PREENCHIDO", "PENDENTE"):
                if bloco is None:
                    raise ErroRedline("sem par")
                if bloco.get("tipo") == "linha_tabela":
                    alvos = celulas_modelo.get(oc["pos_modelo"])
                    if not alvos or len(alvos) != len(bloco["celulas"]):
                        alvos = [t.strip() for t in oc["texto_modelo"].split("|")]
                    reescrever_celulas(bloco["_celulas"], bloco["celulas"], alvos, m)
                else:
                    reescrever_paragrafo(_elemento_bloco(bloco), bloco["texto"],
                                         oc["texto_modelo"], m)
                resumo["revertidas"] += 1
            elif tipo == "INCLUIDO":
                if bloco.get("tipo") == "linha_tabela":
                    excluir_linha(bloco["_tr"], m)
                else:
                    excluir_paragrafo(_elemento_bloco(bloco), m)
                resumo["excluidas"] += 1
            elif tipo in ("REMOVIDO", "MOVIDO"):
                if tipo == "MOVIDO" and bloco is not None:
                    if bloco.get("tipo") == "linha_tabela":
                        excluir_linha(bloco["_tr"], m)
                    else:
                        excluir_paragrafo(_elemento_bloco(bloco), m)
                ancora_pos = _ancora_para_modelo(resultado, oc["pos_modelo"])
                ancora = por_pos.get(ancora_pos)
                celulas = celulas_modelo.get(oc["pos_modelo"])
                if celulas and ancora is not None and ancora.get("_tr") is not None \
                        and ancora.get("tipo") == "linha_tabela":
                    nova_linha_apos(ancora["_tr"], celulas, m)
                else:
                    el = _elemento_bloco(ancora) if ancora else None
                    if el is not None and ancora.get("tipo") == "linha_tabela":
                        el = ancora["_tr"].getparent()          # depois da tabela
                    if el is None:
                        corpo = doc.element.body
                        el = corpo[0] if len(corpo) else None
                        novo = novo_paragrafo_apos(None, oc["texto_modelo"], m)
                        if el is not None:
                            el.addprevious(novo)
                        else:
                            corpo.append(novo)
                    else:
                        novo_paragrafo_apos(el, oc["texto_modelo"], m,
                                            el.find(qn("w:pPr")) if el.tag == qn("w:p")
                                            else None)
                resumo["reinseridas"] += 1
        except (ErroRedline, AttributeError, TypeError) as erro:
            resumo["ignoradas"].append({"id": oc["id"], "motivo": str(erro)[:120]})

    # comentarios da revisao (python-docx >= 1.2)
    if comentarios and hasattr(doc, "add_comment"):
        for oc, decisao in comentarios:
            bloco = por_pos.get(oc["pos_devolucao"] or _ancora_para_modelo(
                resultado, oc["pos_modelo"] or 0))
            el = _elemento_bloco(bloco) if bloco else None
            if el is None:
                continue
            runs = [Run(r, Paragraph(el, doc)) for r in el.iter(qn("w:r"))]
            if not runs:
                continue
            situacao = {"aceito": "ACEITA", "rejeitado": "REJEITADA"}.get(
                decisao.get("status"), "SEM DECISÃO")
            try:
                doc.add_comment(runs[0], text="[%s] %s" % (situacao, decisao["comentario"]),
                                author=autor[:60], initials="".join(
                                    p[0] for p in autor.split()[:2]).upper() or "RV")
                resumo["comentarios"] += 1
            except Exception:  # noqa: BLE001
                pass

    # novas edicoes feitas pelo fornecedor tambem ficam rastreadas
    configuracoes = doc.settings.element
    if configuracoes.find(qn("w:trackRevisions")) is None:
        configuracoes.append(OxmlElement("w:trackRevisions"))

    tmp = caminho_destino + ".tmp"
    doc.save(tmp)
    os.replace(tmp, caminho_destino)
    return resumo


def contar_marcas(caminho):
    """(insercoes, exclusoes, comentarios) de um .docx — usado na verificacao."""
    doc = Document(caminho)
    corpo = doc.element.body
    ins = len(corpo.findall(".//" + qn("w:ins")))
    dele = len(corpo.findall(".//" + qn("w:del")))
    comentarios = len(corpo.findall(".//" + qn("w:commentRangeStart")))
    return ins, dele, comentarios
