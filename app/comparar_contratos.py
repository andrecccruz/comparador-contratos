# -*- coding: utf-8 -*-
"""
comparar_contratos.py
Comparador de devolucoes de contrato x modelo padrao (motor).

Fluxo:  ler_documento_estruturado -> montar_blocos (clausulas, placeholders)
        -> comparar (pareamento + classificacao + risco) -> extrair_ficha
        -> analisar() devolve o resultado estruturado (resultado.json)
        -> relatorios.py gera HTML interativo, PDF, XLSX e TXT a partir dele.

Documentacao: README.md
"""

import os
import re
import sys
import csv
import json
import html
import time
import hashlib
import difflib
import logging
import unicodedata
from datetime import datetime, date

import identidade as identidade_mod

VERSAO = "2.0"
VERSAO_FORMATO_RESULTADO = 2
PASTA_SCRIPT = os.path.dirname(os.path.abspath(__file__))
CAMINHO_CONFIG = os.path.join(PASTA_SCRIPT, "config.json")

TIPOS = ("PREENCHIDO", "ALTERADO", "INCLUIDO", "REMOVIDO", "MOVIDO", "PENDENTE",
         "FORMATACAO")
TIPOS_LEGADO = ("ALTERADO", "INCLUIDO", "REMOVIDO", "MOVIDO")

# ----------------------------------------------------------------------------
# CONFIGURACAO
# ----------------------------------------------------------------------------

TERMOS_CRITICOS_PADRAO = {
    "multa": 15, "rescis": 15, "foro": 12, "pagamento": 12, "pagar": 10,
    "reajuste": 12, "garantia": 10, "confidencial": 10, "lgpd": 12,
    "responsab": 15, "respond": 12, "perdas e danos": 15, "danos": 12,
    "indeniza": 15, "prazo": 8, "vig\u00eancia": 10, "juros": 12, "penalidade": 15,
    "subcontrat": 12, "sigilo": 10, "propriedade intelectual": 12,
    "for\u00e7a maior": 10, "cess\u00e3o": 10, "seguro": 10, "exclusividade": 12,
    "arbitragem": 12, "nota fiscal": 6, "reten\u00e7\u00e3o": 8,
}

CONFIG_PADRAO = {
    "pasta_modelo": "dados/modelo_padr\u00e3o",
    "pasta_devolucoes": "dados/devolucoes",
    "pasta_saida": "dados/sa\u00edda",
    "pasta_logs": "dados/logs",

    "extensoes_aceitas": [".docx", ".pdf", ".txt"],
    "modelo_fixo": "",
    "mapa_modelos": {},

    "formato_data_pasta": "%d.%m.%Y",
    "incluir_hora_na_pasta": True,
    "formato_hora_pasta": "%Hh%M",

    "gerar_html": True,
    "gerar_pdf": True,
    "gerar_txt": True,
    "gerar_xlsx": True,
    "gerar_json": True,
    "incluir_texto_integral": True,       # documento completo no resultado.json/HTML
    "visao_padrao": "inline",             # inline | lado_a_lado | documento
    "copiar_arquivos_origem": False,

    "modo_teste": True,
    "limite_arquivos_teste": 1,
    "gravar_saida": True,

    "similaridade_minima_par": 0.55,
    "similaridade_movido_alterado": 0.80,  # pareamento entre trechos distantes
    "ignorar_diferencas_apenas_formatacao": True,   # FORMATACAO oculta por padrao
    "detectar_blocos_movidos": True,
    "ignorar_maiusculas_minusculas": False,
    "tamanho_minimo_paragrafo": 2,

    "ordem_deteccao_fornecedor": ["subpasta", "conteudo", "arquivo"],
    "separador_nome_arquivo": "__",
    "cnpj_contratante": "",
    "fornecedor_forcado": "",
    # a organizacao contratante (identidade.organizacao_contratante) ja e ignorada
    # automaticamente; liste aqui outros nomes que nunca sao o fornecedor.
    "nomes_ignorar": ["CONTRATANTE"],

    # identidade visual e textual dos relatorios (ver docs/identidade.md)
    "identidade": dict(identidade_mod.IDENTIDADE_PADRAO),
    "caminho_logo": "",

    # termo -> peso no risco (lista simples tambem e aceita: peso 10)
    "termos_criticos": dict(TERMOS_CRITICOS_PADRAO),

    # campos do modelo a preencher (regex); cada ocorrencia vira um "slot"
    "regex_placeholder": [r"\[[^\[\]\n]{1,120}\]", r"_{4,}"],
    "regex_numero_contrato": r"\b[A-Z]{2,4}-[A-Z]{2,5}-\d{4}-\d{1,5}\b",

    "palavras_negacao": ["n\u00e3o", "exceto", "salvo", "vedado", "vedada", "vedados",
                         "proibido", "proibida", "nenhum", "nenhuma", "nunca", "jamais",
                         "ressalvado", "ressalvada"],

    "pesos_risco": {
        "tipo": {"REMOVIDO": 35, "INCLUIDO": 30, "ALTERADO": 25, "PENDENTE": 20,
                 "MOVIDO": 8, "PREENCHIDO": 5, "FORMATACAO": 0},
        "clausula_critica": 15,
        "termo_critico_extra": 3,
        "termo_contexto": 0.6,
        "valor_alterado": 15,
        "valor_presente": 8,
        "negacao": 20,
        "tamanho_max": 15,
    },
    "limiares_risco": {"medio": 30, "alto": 60},
    "categorias_ocultas": ["FORMATACAO"],

    # regras de aceitacao (ver avaliar_playbook); editaveis em Ajustes > Playbook
    "playbook": [],

    "historico_csv": "_historico_comparacoes.csv",
}

# chaves que mudam o resultado da analise (entram no hash de deduplicacao)
CHAVES_RELEVANTES = (
    "similaridade_minima_par", "similaridade_movido_alterado",
    "ignorar_diferencas_apenas_formatacao", "detectar_blocos_movidos",
    "ignorar_maiusculas_minusculas", "tamanho_minimo_paragrafo", "nomes_ignorar",
    "cnpj_contratante", "fornecedor_forcado", "termos_criticos", "regex_placeholder",
    "regex_numero_contrato", "palavras_negacao", "pesos_risco", "limiares_risco",
    "categorias_ocultas", "playbook", "incluir_texto_integral",
)
# a biblioteca so acrescenta um selo informativo; nao muda a analise nem o hash


def carregar_config():
    """Le config.json ao lado do script. Cria template na primeira execucao."""
    if not os.path.exists(CAMINHO_CONFIG):
        with open(CAMINHO_CONFIG, "w", encoding="utf-8") as f:
            json.dump(CONFIG_PADRAO, f, ensure_ascii=False, indent=2)
        print("config.json criado em: %s" % CAMINHO_CONFIG)
        print("Revise os caminhos e execute novamente.")
        sys.exit(0)

    with open(CAMINHO_CONFIG, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    faltantes = [k for k in CONFIG_PADRAO if k not in cfg]
    obsoletas = [k for k in cfg if k not in CONFIG_PADRAO]
    for k in faltantes:
        cfg[k] = CONFIG_PADRAO[k]

    if faltantes:
        print("[CONFIG] Chaves ausentes (usando padrao). Adicione ao config.json:")
        for k in faltantes:
            print("   \"%s\": %s," % (k, json.dumps(CONFIG_PADRAO[k], ensure_ascii=False)))
    if obsoletas:
        print("[CONFIG] Chaves nao reconhecidas (ignoradas): %s" % ", ".join(obsoletas))

    return cfg


def config_relevante(cfg):
    rel = {k: cfg.get(k, CONFIG_PADRAO.get(k)) for k in CHAVES_RELEVANTES}
    rel["organizacao_contratante"] = (cfg.get("identidade") or {}).get(
        "organizacao_contratante", "")
    return rel


def _sha256_arquivo(h, caminho):
    with open(caminho, "rb") as f:
        for pedaco in iter(lambda: f.read(1 << 16), b""):
            h.update(pedaco)


def hash_entrada(caminho_modelo, caminho_devolucao, cfg, conteudo_devolucao=None):
    """SHA-256 de (modelo + devolucao + versao do motor + config relevante).
    `conteudo_devolucao` (bytes) permite calcular antes de gravar o upload."""
    h = hashlib.sha256()
    _sha256_arquivo(h, caminho_modelo)
    h.update(b"\x00")
    if conteudo_devolucao is not None:
        h.update(conteudo_devolucao)
    else:
        _sha256_arquivo(h, caminho_devolucao)
    h.update(b"\x00" + VERSAO.encode("ascii") + b"\x00")
    h.update(json.dumps(config_relevante(cfg), sort_keys=True,
                        ensure_ascii=False).encode("utf-8"))
    return h.hexdigest()


# ----------------------------------------------------------------------------
# UTILITARIOS DE TEXTO
# ----------------------------------------------------------------------------

NBSP = "\u00a0"
_TRADUCOES = {
    NBSP: " ", "\u2013": "-", "\u2014": "-", "\u2212": "-",
    "\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201c": '"',
    "\u201d": '"', "\u201e": '"', "\u2026": "...", "\u00ad": "",
    "\ufeff": "", "\t": " ",
}


def normalizar(texto, ignorar_caixa=False):
    """Normaliza para COMPARACAO (nao para exibicao)."""
    t = unicodedata.normalize("NFKC", texto or "")
    for k, v in _TRADUCOES.items():
        t = t.replace(k, v)
    t = re.sub(r"\s+", " ", t).strip()
    t = re.sub(r"\s+([,.;:!?%\)\]])", r"\1", t)
    t = re.sub(r"([\(\[])\s+", r"\1", t)
    if ignorar_caixa:
        t = t.lower()
    return t


def normalizar_leve(texto):
    """So espacos e aspas/travessoes: preserva o texto como o usuario escreveu
    (usado para extrair valores preenchidos)."""
    t = texto or ""
    for k, v in _TRADUCOES.items():
        t = t.replace(k, v)
    return re.sub(r"\s+", " ", t).strip()


def sem_acento(texto):
    t = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in t if not unicodedata.combining(c))


def chave_formatacao(texto):
    """Chave que ignora espacos, pontuacao e caixa (categoria FORMATACAO)."""
    t = normalizar(texto).lower()
    return re.sub(r"[\s\.,;:!?\"'()\[\]\-/]+", "", t)


def sanitizar_nome_pasta(nome, limite=70):
    nome = re.sub(r"[<>:\"/\\|?*\r\n\t]", " ", nome or "")
    nome = re.sub(r"\s+", "_", nome.strip())
    nome = nome.strip("._-")
    if not nome:
        nome = "FORNECEDOR_NAO_IDENTIFICADO"
    return nome[:limite]


def proporcao_maiusculas(texto):
    letras = [c for c in texto if c.isalpha()]
    if not letras:
        return 0.0
    return sum(1 for c in letras if c.isupper()) / float(len(letras))


def validar_cnpj(cnpj):
    """Valida os digitos verificadores de um CNPJ (com ou sem mascara)."""
    numeros = re.sub(r"\D", "", cnpj or "")
    if len(numeros) != 14 or numeros == numeros[0] * 14:
        return False

    def digito(base, pesos):
        soma = sum(int(n) * p for n, p in zip(base, pesos))
        resto = soma % 11
        return "0" if resto < 2 else str(11 - resto)

    pesos1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    pesos2 = [6] + pesos1
    d1 = digito(numeros[:12], pesos1)
    d2 = digito(numeros[:12] + d1, pesos2)
    return numeros[-2:] == d1 + d2


def _termos_com_peso(cfg):
    termos = cfg.get("termos_criticos", TERMOS_CRITICOS_PADRAO)
    if isinstance(termos, dict):
        return {str(k): float(v) for k, v in termos.items() if str(k).strip()}
    return {str(t): 10.0 for t in (termos or []) if str(t).strip()}


def _regexes_placeholder(cfg):
    saida = []
    for padrao in cfg.get("regex_placeholder") or CONFIG_PADRAO["regex_placeholder"]:
        try:
            saida.append(re.compile(padrao))
        except re.error:
            continue
    return saida


def encontrar_placeholders(texto, regexes):
    """Spans (inicio, fim) nao sobrepostos dos campos a preencher."""
    spans = []
    for regex in regexes:
        for m in regex.finditer(texto):
            if m.end() > m.start():
                spans.append((m.start(), m.end()))
    spans.sort()
    unidos = []
    for s, e in spans:
        if unidos and s <= unidos[-1][1]:
            unidos[-1] = (unidos[-1][0], max(e, unidos[-1][1]))
        elif (unidos and texto[s] == "_" and texto[unidos[-1][1] - 1] == "_"
              and re.match(r"^[\s/.\-]{1,3}$", texto[unidos[-1][1]:s])):
            unidos[-1] = (unidos[-1][0], e)        # ____/____/____ = um campo so
        else:
            unidos.append((s, e))
    return unidos


def rotulo_campo(texto, inicio, fim):
    """Nome legivel do campo: '[NOME]' ou, para linhas em branco, 'Data: ____'."""
    campo = texto[inicio:fim]
    if not campo.startswith("_"):
        return campo
    antes = re.search(r"([A-Za-z\u00c0-\u00ff\. ]{2,25}):?\s*$",
                      texto[max(0, inicio - 30):inicio])
    return ("%s: ____" % antes.group(1).strip()) if antes else "____"


# ----------------------------------------------------------------------------
# LEITURA DE ARQUIVOS
# ----------------------------------------------------------------------------

def _nivel_titulo(paragrafo):
    """Nivel de titulo pelo estilo (Heading N / Titulo N) ou outlineLvl."""
    from docx.oxml.ns import qn
    estilo = paragrafo.style
    while estilo is not None:
        nome = (estilo.name or "").lower()
        m = re.match(r"^(heading|t[i\u00ed]tulo)\s*(\d)", nome)
        if m:
            return int(m.group(2))
        if nome in ("title", "t\u00edtulo", "titulo"):
            return 0
        estilo = estilo.base_style
    ppr = paragrafo._p.pPr
    if ppr is not None:
        nivel = ppr.find(qn("w:outlineLvl"))
        if nivel is not None:
            try:
                return int(nivel.get(qn("w:val"))) + 1
            except (TypeError, ValueError):
                return None
    return None


def _todo_negrito(paragrafo):
    runs = [r for r in paragrafo.runs if r.text.strip()]
    if not runs:
        return False
    estilo_negrito = bool(paragrafo.style is not None and paragrafo.style.font.bold)
    return all(r.bold or (r.bold is None and estilo_negrito) for r in runs)


def ler_docx_estruturado(caminho, com_elementos=False, documento=None):
    """Blocos do .docx. Com `com_elementos`, cada bloco leva `_el` (w:p) ou
    `_celulas` (w:tc da linha) do `documento` informado — usado pelo redline."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    doc = documento if documento is not None else Document(caminho)
    blocos = []
    n_tabela = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            par = Paragraph(child, doc)
            texto = (par.text or "").strip()
            if texto:
                blocos.append({"texto": texto, "tipo": "paragrafo",
                               "estilo": par.style.name if par.style is not None else "",
                               "negrito": _todo_negrito(par),
                               "nivel_titulo": _nivel_titulo(par)})
                if com_elementos:
                    blocos[-1]["_el"] = child
        elif child.tag == qn("w:tbl"):
            tabela = Table(child, doc)
            n_tabela += 1
            linhas, elementos = [], []
            for linha in tabela.rows:
                celulas, tcs, vistos = [], [], set()
                for c in linha.cells:
                    if id(c._tc) in vistos:
                        continue
                    vistos.add(id(c._tc))
                    celulas.append(re.sub(r"\s+", " ", c.text).strip())
                    tcs.append(c._tc)
                linhas.append(celulas)
                elementos.append((linha._tr, tcs))
            nao_vazias = [[c for c in l if c] for l in linhas]
            if all(len(l) <= 1 for l in nao_vazias):
                # tabela de uma coluna (ex.: titulo dentro de tabela): vira paragrafo
                for l, (tr, tcs) in zip(nao_vazias, elementos):
                    if l:
                        blocos.append({"texto": l[0], "tipo": "paragrafo",
                                       "estilo": "tabela", "negrito": False,
                                       "nivel_titulo": None, "em_tabela": True})
                        if com_elementos:
                            blocos[-1]["_celulas"] = tcs
                            blocos[-1]["_tr"] = tr
                continue
            for i, celulas in enumerate(linhas):
                conteudo = " | ".join(c for c in celulas if c)
                if conteudo.strip(" |"):
                    blocos.append({"texto": conteudo, "tipo": "linha_tabela",
                                   "tabela": n_tabela, "linha": i,
                                   "colunas": len(celulas), "celulas": celulas})
                    if com_elementos:
                        blocos[-1]["_tr"], blocos[-1]["_celulas"] = elementos[i]
    return blocos


def ler_pdf(caminho):
    texto_total = []
    try:
        import pdfplumber
        with pdfplumber.open(caminho) as pdf:
            for pagina in pdf.pages:
                texto_total.append(pagina.extract_text() or "")
    except ImportError:
        from pypdf import PdfReader
        leitor = PdfReader(caminho)
        for pagina in leitor.pages:
            texto_total.append(pagina.extract_text() or "")

    linhas = []
    for pagina in texto_total:
        for linha in pagina.splitlines():
            linhas.append(linha.strip())

    # remonta paragrafos: junta linhas que nao terminam em pontuacao forte;
    # linhas de titulo (CLAUSULA..., 1. TITULO) ficam sozinhas
    blocos = []
    buffer_ = ""
    for linha in linhas:
        if not linha:
            if buffer_:
                blocos.append(buffer_.strip())
                buffer_ = ""
            continue
        if re.match(r"^\s*(P[\u00c1A]GINA|Page)\s+\d+", linha, re.IGNORECASE):
            continue
        if _parece_titulo_texto(linha):
            if buffer_:
                blocos.append(buffer_.strip())
                buffer_ = ""
            blocos.append(linha)
            continue
        buffer_ = (buffer_ + " " + linha).strip() if buffer_ else linha
        if re.search(r"[.;:!?]\s*$", linha):
            blocos.append(buffer_.strip())
            buffer_ = ""
    if buffer_:
        blocos.append(buffer_.strip())
    return [b for b in blocos if b.strip()]


def ler_txt(caminho):
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with open(caminho, "r", encoding=enc) as f:
                conteudo = f.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        raise IOError("Nao foi possivel decodificar %s" % caminho)

    blocos = []
    for bloco in re.split(r"\n\s*\n", conteudo):
        for linha in bloco.splitlines():
            if linha.strip():
                blocos.append(linha.strip())
    return blocos


def ler_documento_estruturado(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    if ext == ".docx":
        return ler_docx_estruturado(caminho)
    if ext == ".pdf":
        textos = ler_pdf(caminho)
    elif ext in (".txt", ".md"):
        textos = ler_txt(caminho)
    else:
        raise ValueError("Extensao nao suportada: %s (converta .doc para .docx)" % ext)
    return [{"texto": t, "tipo": "paragrafo", "estilo": "", "negrito": False,
             "nivel_titulo": None} for t in textos]


def ler_documento(caminho):
    """Compatibilidade: lista de textos (usada na leitura de modelos na web)."""
    return [b["texto"] for b in ler_documento_estruturado(caminho)]


# ----------------------------------------------------------------------------
# ESTRUTURA / CLAUSULAS
# ----------------------------------------------------------------------------

_MAIUSC = "A-Z\u00c0-\u00dd"
RE_CLAUSULA = re.compile(r"^\s*CL[\u00c1A]USULA\s+\S+", re.IGNORECASE)
RE_ANEXO = re.compile(r"^\s*ANEXO\s+([IVXLCDM]+|\d{1,3}|[A-Z])\b(?!\w)", re.IGNORECASE)
RE_TITULO_NUMERADO = re.compile(
    r"^\s*(\d{1,2}|[IVXLC]{1,6})\s*[\.\)\-\u2013\u2014]?\s+([%s].{1,110})$" % _MAIUSC)
RE_PARAGRAFO = re.compile(
    r"^\s*(PAR[\u00c1A]GRAFO\s+[\u00daU]NICO|PAR[\u00c1A]GRAFO\s+\w+|\u00a7\s*\d+[\u00ba\u00b0]?)",
    re.IGNORECASE)
RE_ITEM = re.compile(r"^\s*(\d{1,3}(?:\.\d{1,3}){1,4})[\.\)\-\s]")
RE_SEPARADOR_TITULO = re.compile(r"\s*(?::|\.\s|\s[-\u2013\u2014]\s)")
PREAMBULO = "PRE\u00c2MBULO"


def _limpar_titulo(texto, limite=110):
    limpo = re.sub(r"\s+", " ", texto).strip(" .;:")
    if len(limpo) > limite:
        # titulo seguido do corpo no mesmo paragrafo: corta no 2o separador
        cortes = [m.start() for m in RE_SEPARADOR_TITULO.finditer(limpo)]
        corte = cortes[1] if len(cortes) > 1 else (cortes[0] if cortes else 80)
        limpo = limpo[:min(corte, limite)].rstrip(" .;:-\u2013")
    return limpo.upper()


def _parece_titulo_texto(texto):
    """Titulo reconhecivel so pelo texto (usado tambem na leitura de PDF)."""
    t = texto.strip()
    if RE_CLAUSULA.match(t) and len(t) <= 160:
        return True
    if RE_ANEXO.match(t) and len(t) <= 160:
        return True
    m = RE_TITULO_NUMERADO.match(t)
    if m and len(t) <= 110 and not t.endswith((";", ",")):
        return proporcao_maiusculas(m.group(2)) >= 0.7
    return False


def classificar_titulo(bloco):
    """Devolve ('forte'|'fraco', rotulo) quando o bloco e titulo de clausula."""
    texto = bloco["texto"].strip()
    if bloco.get("tipo") == "linha_tabela":
        return None
    if RE_CLAUSULA.match(texto) or RE_ANEXO.match(texto):
        return ("forte", _limpar_titulo(texto))
    if _parece_titulo_texto(texto):
        return ("forte", _limpar_titulo(texto))
    nivel = bloco.get("nivel_titulo")
    curto = len(texto) <= 120 and not texto.endswith((";", ","))
    if nivel is not None and 1 <= nivel <= 2 and curto:
        return ("forte", _limpar_titulo(texto))
    if (bloco.get("negrito") and len(texto) <= 90 and proporcao_maiusculas(texto) >= 0.7
            and not texto.endswith((".", ";", ":", ","))):
        return ("fraco", _limpar_titulo(texto))
    return None


def rotulo_estrutural(texto):
    """Rotulo curto do paragrafo (compatibilidade): clausula, paragrafo ou item."""
    titulo = classificar_titulo({"texto": texto, "tipo": "paragrafo"})
    if titulo:
        return titulo[1]
    m = RE_PARAGRAFO.match(texto)
    if m:
        return re.sub(r"\s+", " ", m.group(1)).strip(" .-").upper()
    m = RE_ITEM.match(texto)
    if m:
        return "Item %s" % m.group(1)
    return None


def montar_blocos(paragrafos, cfg):
    """Lista de textos ou blocos estruturados -> blocos com clausula, norm e
    placeholders. Titulos 'fracos' (negrito em caixa alta) so abrem clausula
    depois do primeiro titulo 'forte' (evita que o titulo do documento vire
    clausula)."""
    blocos = []
    clausula_atual = PREAMBULO
    ordem_clausula = 0
    viu_forte = False
    minimo = int(cfg.get("tamanho_minimo_paragrafo", 2))
    ignorar_caixa = cfg.get("ignorar_maiusculas_minusculas", False)
    regexes = _regexes_placeholder(cfg)

    for item in paragrafos:
        bloco = dict(item) if isinstance(item, dict) else {"texto": item,
                                                           "tipo": "paragrafo"}
        texto = bloco["texto"]
        titulo = classificar_titulo(bloco)
        eh_titulo = False
        if titulo and (titulo[0] == "forte" or viu_forte):
            viu_forte = viu_forte or titulo[0] == "forte"
            clausula_atual = titulo[1]
            ordem_clausula += 1
            eh_titulo = True
        norm = normalizar(texto, ignorar_caixa)
        if len(norm) < minimo:
            continue
        bloco.update({
            "norm": norm,
            "clausula": clausula_atual,
            "ordem_clausula": ordem_clausula,
            "titulo": eh_titulo,
            "rotulo": rotulo_estrutural(texto) or "",
            "placeholders": [rotulo_campo(norm, s, e)
                             for s, e in encontrar_placeholders(norm, regexes)],
        })
        blocos.append(bloco)
    for pos, bloco in enumerate(blocos, 1):
        bloco["pos"] = pos
    return blocos


# ----------------------------------------------------------------------------
# IDENTIFICACAO DO FORNECEDOR
# ----------------------------------------------------------------------------

RE_CNPJ = re.compile(r"\b(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})\b")

# nivel 1: nome que termina em forma societaria (greedy -> pega o nome completo)
RE_RAZAO_FORTE = re.compile(
    r"([A-Z\u00c0-\u00dc0-9][A-Z\u00c0-\u00dc0-9&\.\'\- ]{3,90}"
    r"\s(?:LTDA|S\.\s?A\.|S/A|EIRELI|EPP|MEI|SOCIEDADE AN\u00d4NIMA)\b\.?)")
# nivel 2: nome com termo de atividade, usado so quando o nivel 1 nao acha nada
RE_RAZAO_FRACA = re.compile(
    r"([A-Z\u00c0-\u00dc0-9][A-Z\u00c0-\u00dc0-9&\.\'\- ]{3,90}"
    r"\s(?:SERVI\u00c7OS|SERVICOS|ENGENHARIA|CONSTRU\u00c7\u00d5ES|CONSTRUTORA|"
    r"COM\u00c9RCIO|COMERCIO|IND\u00daSTRIA|INDUSTRIA|TECNOLOGIA|"
    r"EMPREENDIMENTOS|PARTICIPA\u00c7\u00d5ES|ASSESSORIA|CONSULTORIA)\b\.?)")
RE_PREFIXO_PAPEL = re.compile(
    r"^(?:A|O|AS|OS|EMPRESA|FORNECEDOR[A]?|CONTRATAD[AO]|CONTRATANTE|"
    r"PRESTADOR[A]?|RAZ\u00c3O SOCIAL)\b[\s:\-,]*", re.IGNORECASE)


def _nomes_bloqueados(cfg):
    nomes = list(cfg.get("nomes_ignorar", []))
    organizacao = (cfg.get("identidade") or {}).get("organizacao_contratante", "")
    if organizacao and len(organizacao.strip()) >= 4:
        nomes.append(organizacao.strip())
    return nomes


def _nome_bloqueado(nome, cfg):
    alvo = sem_acento(nome).upper()
    for proibido in _nomes_bloqueados(cfg):
        if sem_acento(proibido).upper() in alvo:
            return True
    return False


def detectar_por_conteudo(blocos, cfg):
    """Procura razao social + CNPJ na qualificacao das partes."""
    trecho = " \n ".join(b["texto"] for b in blocos[:80])
    cnpj_contratante = (cfg.get("cnpj_contratante") or "").strip()

    def coletar(regex, bonus):
        achados = []
        for linha in trecho.split("\n"):
            alto = unicodedata.normalize("NFKC", linha).upper()
            for m in regex.finditer(alto):
                nome = RE_PREFIXO_PAPEL.sub("", m.group(1))
                nome = re.sub(r"\s+", " ", nome).strip(" ,.-")
                if len(nome) < 6 or _nome_bloqueado(nome, cfg):
                    continue
                achado_cnpj = RE_CNPJ.search(linha)
                cnpj = achado_cnpj.group(1) if achado_cnpj else ""
                if cnpj and cnpj_contratante and cnpj == cnpj_contratante:
                    continue
                peso = bonus + (2 if cnpj else 0)
                if re.search(r"CONTRATAD[AO]", alto):
                    peso += 3
                achados.append((peso, nome, cnpj))
        return achados

    candidatos = coletar(RE_RAZAO_FORTE, 5) or coletar(RE_RAZAO_FRACA, 1)

    # entre nomes sobrepostos, mantem o mais completo
    candidatos = [c for c in candidatos
                  if not any(c[1] != o[1] and c[1] in o[1] for o in candidatos)]

    if not candidatos:
        m = RE_CNPJ.search(trecho)
        return ("", m.group(1) if m else "")

    candidatos.sort(key=lambda c: (-c[0], -len(c[1])))
    return (candidatos[0][1], candidatos[0][2])


def detectar_fornecedor(caminho_arquivo, cfg, blocos):
    nome, cnpj = "", ""

    forcado = (cfg.get("fornecedor_forcado") or "").strip()
    if forcado:                       # informado manualmente (ex.: pela interface)
        _, cnpj = detectar_por_conteudo(blocos, cfg)
        return (forcado, cnpj)

    raiz_devolucoes = os.path.normpath(cfg.get("pasta_devolucoes") or "")

    for origem in cfg.get("ordem_deteccao_fornecedor", []):
        if nome:
            break
        if origem == "subpasta" and raiz_devolucoes:
            pasta = os.path.dirname(os.path.normpath(caminho_arquivo))
            if pasta != raiz_devolucoes and pasta.startswith(raiz_devolucoes):
                relativo = os.path.relpath(pasta, raiz_devolucoes)
                candidato = relativo.split(os.sep)[0]
                if candidato and candidato not in (".", ".."):
                    nome = candidato.replace("_", " ").strip()
        elif origem == "arquivo":
            base = os.path.splitext(os.path.basename(caminho_arquivo))[0]
            sep = cfg.get("separador_nome_arquivo") or "__"
            if sep in base:
                nome = base.split(sep)[0].replace("_", " ").strip()
        elif origem == "conteudo":
            n, c = detectar_por_conteudo(blocos, cfg)
            if n:
                nome = n
            if c:
                cnpj = c

    if not cnpj:
        _, cnpj = detectar_por_conteudo(blocos, cfg)

    return (nome.strip() or "FORNECEDOR NAO IDENTIFICADO", cnpj)


# ----------------------------------------------------------------------------
# DIFF EM NIVEL DE PALAVRA
# ----------------------------------------------------------------------------

RE_TOKEN = re.compile(r"""
    \[[^\[\]\n]{1,120}\]                          # campo do modelo
  | R\$\s?\d[\d.]*(?:,\d+)?                        # valor em reais
  | \d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}               # CNPJ
  | \d{3}\.\d{3}\.\d{3}-\d{2}                     # CPF
  | \d{1,2}/\d{1,2}/\d{2,4}                       # data
  | \d+(?:[.,]\d+)*\s?%                           # percentual
  | \d+(?:[.,]\d+)*[\u00ba\u00aa\u00b0]?           # numero / ordinal
  | \w+(?:[-']\w+)*                               # palavra (com hifen)
  | \s+
  | [^\w\s]
""", re.VERBOSE | re.UNICODE)


def tokenizar(texto):
    return RE_TOKEN.findall(texto or "")


def diff_ops(texto_a, texto_b):
    """Diff por token -> lista [op, texto] com op em '=', '-', '+'.
    Trechos iguais curtos entre duas mudancas sao absorvidos (diff mais limpo)."""
    ta, tb = tokenizar(texto_a), tokenizar(texto_b)
    sm = difflib.SequenceMatcher(None, ta, tb, autojunk=False)
    brutos = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        sa, sb = "".join(ta[i1:i2]), "".join(tb[j1:j2])
        if tag == "equal":
            brutos.append(["=", sa, sb])
        else:
            brutos.append(["~", sa, sb])
    # absorve "iguais" irrelevantes (espaco, pontuacao, ate 2 caracteres) entre mudancas
    limpos = []
    for i, (op, sa, sb) in enumerate(brutos):
        entre = (0 < i < len(brutos) - 1 and op == "="
                 and brutos[i - 1][0] == "~" and brutos[i + 1][0] == "~"
                 and (not sa.strip() or len(sa.strip()) <= 2))
        if entre:
            op = "~"
        if limpos and limpos[-1][0] == "~" and op == "~":
            limpos[-1][1] += sa
            limpos[-1][2] += sb
        else:
            limpos.append([op, sa, sb])
    ops = []
    for op, sa, sb in limpos:
        if op == "=":
            ops.append(["=", sa])
        else:
            if sa:
                ops.append(["-", sa])
            if sb:
                ops.append(["+", sb])
    return ops


def html_de_ops(ops):
    """(html_modelo, html_devolucao, marcado) a partir das ops."""
    lado_a, lado_b, marcado = [], [], []
    for op, texto in ops:
        e = html.escape(texto)
        if op == "=":
            lado_a.append(e)
            lado_b.append(e)
            marcado.append(texto)
        elif op == "-":
            lado_a.append('<span class="del">%s</span>' % e)
            marcado.append("[--%s--]" % texto.strip())
        else:
            lado_b.append('<span class="ins">%s</span>' % e)
            marcado.append("{++%s++}" % texto.strip())
    return "".join(lado_a), "".join(lado_b), "".join(marcado)


def diff_inline(texto_a, texto_b):
    """Compatibilidade v1: (html_original, html_novo, texto_marcado)."""
    return html_de_ops(diff_ops(texto_a, texto_b))


# ----------------------------------------------------------------------------
# CAMPOS DO MODELO (PREENCHIDO / PENDENTE)
# ----------------------------------------------------------------------------

def _casar_texto_modelo(texto_modelo, texto_dev, regexes):
    """Compara o texto do modelo (com campos) contra a devolucao ancorando os trechos
    literais em ordem (linear, sem backtracking). Devolve lista de
    (placeholder, valor) ou None se o texto ao redor mudou."""
    spans = encontrar_placeholders(texto_modelo, regexes)
    if not spans:
        return None
    literais, campos = [], []
    ultimo = 0
    for s, e in spans:
        literais.append(texto_modelo[ultimo:s])
        campos.append(rotulo_campo(texto_modelo, s, e))
        ultimo = e
    literais.append(texto_modelo[ultimo:])

    alvo = texto_dev.lower()
    lits = [l.lower() for l in literais]
    if not alvo.startswith(lits[0]):
        return None
    pos = len(lits[0])
    valores = []
    for k in range(len(campos)):
        proximo = lits[k + 1]
        if k + 1 == len(campos):                   # ultimo literal: precisa ser sufixo
            if proximo and not alvo.endswith(proximo):
                return None
            fim = len(alvo) - len(proximo)
            if fim < pos:
                return None
        elif proximo:
            fim = alvo.find(proximo, pos + 1)
            if fim < 0:
                return None
        else:                                      # campos colados: valor vai no 1o
            fim = pos
        valores.append(texto_dev[pos:fim].strip())
        pos = fim + len(proximo)
    return list(zip(campos, valores))


def casar_modelo(a, b, regexes):
    """Resultado do casamento modelo x devolucao para um par de blocos:
    {'campos': [{'placeholder','valor'}], 'pendentes': [...]} ou None."""
    if not a.get("placeholders"):
        return None
    pares = None
    if (a.get("tipo") == "linha_tabela" and b.get("tipo") == "linha_tabela"
            and a.get("colunas") == b.get("colunas")):
        pares = []
        for ca, cb in zip(a["celulas"], b["celulas"]):
            na, nb = normalizar_leve(ca), normalizar_leve(cb)
            if encontrar_placeholders(na, regexes):
                casado = _casar_texto_modelo(na, nb, regexes)
                if casado is None:
                    return None
                pares.extend(casado)
            elif chave_formatacao(na) != chave_formatacao(nb):
                return None                         # coluna fixa mudou
    else:
        pares = _casar_texto_modelo(normalizar_leve(a["texto"]),
                                    normalizar_leve(b["texto"]), regexes)
        if pares is None:
            return None
    campos, pendentes = [], []
    for placeholder, valor in pares:
        if not valor or encontrar_placeholders(valor, regexes):
            pendentes.append(placeholder)
        else:
            campos.append({"placeholder": placeholder, "valor": valor})
    return {"campos": campos, "pendentes": pendentes}


# ----------------------------------------------------------------------------
# PAREAMENTO
# ----------------------------------------------------------------------------

def _razao(a, b):
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    if sm.real_quick_ratio() < 0.3 or sm.quick_ratio() < 0.3:
        return 0.0
    return sm.ratio()


def _colunas_fixas_iguais(a, b, regexes):
    """Linhas de tabela com a mesma estrutura: colunas sem campo sao iguais."""
    if a.get("tipo") != "linha_tabela" or b.get("tipo") != "linha_tabela":
        return False
    if a.get("colunas") != b.get("colunas"):
        return False
    fixas = 0
    for ca, cb in zip(a["celulas"], b["celulas"]):
        if encontrar_placeholders(ca, regexes):
            continue
        if chave_formatacao(ca) != chave_formatacao(cb):
            return False
        if ca.strip():
            fixas += 1
    return fixas > 0 or a.get("linha") == b.get("linha")


def _parear_faixa(ia, ib, A, B, cfg, regexes, distante=False):
    """Pareia indices de A e B que nao ficaram 'iguais' no alinhamento.
    Ordem de evidencia: texto identico (movido) > campos do modelo > estrutura
    de tabela > similaridade. Devolve (pares, sobra_a, sobra_b)."""
    livres_a, livres_b = list(ia), list(ib)
    pares = []

    def usar(i, j, motivo, razao):
        pares.append((i, j, motivo, razao))
        livres_a.remove(i)
        livres_b.remove(j)

    # 1) texto identico em outra posicao
    por_norm = {}
    for j in livres_b:
        por_norm.setdefault(B[j]["norm"], []).append(j)
    for i in list(livres_a):
        candidatos = [j for j in por_norm.get(A[i]["norm"], []) if j in livres_b]
        if candidatos:
            usar(i, candidatos[0], "movido", 1.0)

    # 2) campos do modelo preenchidos com o texto ao redor intacto
    for i in list(livres_a):
        if not A[i].get("placeholders"):
            continue
        literal = re.sub(r"\s", "", "".join(
            A[i]["norm"][ultimo:s] for ultimo, s in _intervalos_literais(A[i], regexes)))
        if len(literal) < 3 and A[i].get("tipo") != "linha_tabela":
            continue                                # so campo, sem ancora: arriscado
        melhores = []
        for j in livres_b:
            if casar_modelo(A[i], B[j], regexes) is not None:
                melhores.append((abs((j - ib[0]) - (i - ia[0])) if ib and ia else 0, j))
        if melhores:
            melhores.sort()
            usar(i, melhores[0][1], "distante" if distante else "modelo", 1.0)

    # 3) linhas de tabela com a mesma estrutura
    for i in list(livres_a):
        if A[i].get("tipo") != "linha_tabela":
            continue
        for j in livres_b:
            if _colunas_fixas_iguais(A[i], B[j], regexes):
                usar(i, j, "tabela", _razao(A[i]["norm"], B[j]["norm"]))
                break

    # 4) similaridade (gulosa, maior razao primeiro)
    minimo = float(cfg.get("similaridade_movido_alterado", 0.8) if distante
                   else cfg.get("similaridade_minima_par", 0.55))
    candidatos = []
    for i in livres_a:
        for j in livres_b:
            r = _razao(A[i]["norm"], B[j]["norm"])
            if r >= minimo:
                candidatos.append((r, -abs(i - j), i, j))
    candidatos.sort(reverse=True)
    for r, _, i, j in candidatos:
        if i in livres_a and j in livres_b:
            usar(i, j, "distante" if distante else "similar", r)

    return pares, livres_a, livres_b


def _intervalos_literais(bloco, regexes):
    """Pares (inicio, fim) dos trechos literais (fora dos campos) do texto norm."""
    spans = encontrar_placeholders(bloco["norm"], regexes)
    saida, ultimo = [], 0
    for s, e in spans:
        saida.append((ultimo, s))
        ultimo = e
    saida.append((ultimo, len(bloco["norm"])))
    return saida


# ----------------------------------------------------------------------------
# RISCO
# ----------------------------------------------------------------------------

RE_VALOR = re.compile(
    r"R\$\s?\d[\d.]*(?:,\d+)?"
    r"|\d+(?:[.,]\d+)*\s?%"
    r"|\d{1,2}/\d{1,2}/\d{2,4}"
    r"|\d+(?:[.,]\d+)*\s*(?:\([^)]{1,40}\)\s*)?"
    r"(?:dias?(?:\s+(?:[u\u00fa]teis|[u\u00fa]til|corridos))?|horas?|h\b|meses|m[e\u00ea]s|anos?|semanas?)",
    re.IGNORECASE)


def valores_do_texto(texto):
    return [re.sub(r"\s+", " ", v).strip().lower() for v in RE_VALOR.findall(texto or "")]


def _selo(risco, cfg):
    limiares = cfg.get("limiares_risco") or CONFIG_PADRAO["limiares_risco"]
    if risco >= float(limiares.get("alto", 60)):
        return "Alto"
    if risco >= float(limiares.get("medio", 30)):
        return "M\u00e9dio"
    return "Baixo"


def _contar_negacoes(texto, palavras):
    alvo = " %s " % sem_acento(texto or "").lower()
    total = {}
    for p in palavras:
        chave = sem_acento(p).lower()
        n = len(re.findall(r"(?<!\w)%s(?!\w)" % re.escape(chave), alvo))
        if n:
            total[p] = n
    return total


def avaliar_risco(oc, cfg, clausula_critica):
    """Nivel de risco 0-100 e fatores explicativos de uma ocorrencia."""
    pesos = dict(CONFIG_PADRAO["pesos_risco"])
    pesos.update(cfg.get("pesos_risco") or {})
    pesos_tipo = dict(CONFIG_PADRAO["pesos_risco"]["tipo"])
    pesos_tipo.update((cfg.get("pesos_risco") or {}).get("tipo") or {})
    termos = _termos_com_peso(cfg)
    tipo = oc["tipo"]
    fatores = []

    risco = float(pesos_tipo.get(tipo, 0))
    fatores.append("%s (+%d)" % (tipo, risco))
    if tipo == "FORMATACAO":
        return 0, fatores, []

    if tipo in ("ALTERADO", "PREENCHIDO", "PENDENTE"):
        trecho = " ".join(t for op, t in oc.get("diff", []) if op in "-+")
        if tipo == "PENDENTE":
            trecho = oc.get("texto_devolucao", "")
    else:
        trecho = oc.get("texto_modelo", "") + " " + oc.get("texto_devolucao", "")
    alvo = sem_acento(trecho).lower()

    atingidos = sorted(((p, t) for t, p in termos.items()
                        if sem_acento(t).lower() in alvo), reverse=True)
    termos_hit = [t for _, t in atingidos]
    if atingidos and tipo not in ("PREENCHIDO",):
        risco += atingidos[0][0]
        extra = min(len(atingidos) - 1, 3) * float(pesos.get("termo_critico_extra", 3))
        risco += extra
        fatores.append("termo cr\u00edtico: %s (+%d)"
                       % (", ".join(termos_hit[:4]), atingidos[0][0] + extra))
    elif tipo == "ALTERADO":
        # termo no restante do paragrafo (mudanca pontual em frase critica)
        contexto = sem_acento(oc.get("texto_devolucao", "")).lower()
        no_contexto = sorted(((p, t) for t, p in termos.items()
                              if sem_acento(t).lower() in contexto), reverse=True)
        if no_contexto:
            fator = float(pesos.get("termo_contexto", 0.6))
            ganho = no_contexto[0][0] * fator
            risco += ganho
            termos_hit = [t for _, t in no_contexto]
            fatores.append("termo cr\u00edtico no contexto: %s (+%d)"
                           % (", ".join(termos_hit[:4]), ganho))
    if clausula_critica and tipo not in ("PREENCHIDO",):
        risco += float(pesos.get("clausula_critica", 15))
        fatores.append("cl\u00e1usula cr\u00edtica (+%d)" % pesos.get("clausula_critica", 15))

    va = valores_do_texto(oc.get("texto_modelo", ""))
    vb = valores_do_texto(oc.get("texto_devolucao", ""))
    if tipo in ("ALTERADO", "MOVIDO") and sorted(va) != sorted(vb):
        risco += float(pesos.get("valor_alterado", 15))
        saiu = [v for v in va if v not in vb]
        entrou = [v for v in vb if v not in va]
        fatores.append("valor alterado: %s \u2192 %s (+%d)" % (
            ", ".join(saiu[:3]) or "-", ", ".join(entrou[:3]) or "-",
            pesos.get("valor_alterado", 15)))
    elif tipo in ("INCLUIDO", "REMOVIDO") and (va or vb):
        risco += float(pesos.get("valor_presente", 8))
        fatores.append("cont\u00e9m valores/prazos (+%d)" % pesos.get("valor_presente", 8))

    if tipo not in ("PREENCHIDO", "PENDENTE"):
        palavras = cfg.get("palavras_negacao") or CONFIG_PADRAO["palavras_negacao"]
        na = _contar_negacoes(oc.get("texto_modelo", ""), palavras)
        nb = _contar_negacoes(oc.get("texto_devolucao", ""), palavras)
        if tipo in ("INCLUIDO", "REMOVIDO"):
            mudou = sorted(set(na) | set(nb))
        else:
            mudou = sorted(p for p in set(na) | set(nb) if na.get(p, 0) != nb.get(p, 0))
        if mudou:
            risco += float(pesos.get("negacao", 20))
            sentido = ("introduzida" if sum(nb.values()) > sum(na.values())
                       else "removida")
            fatores.append("nega\u00e7\u00e3o/exce\u00e7\u00e3o %s: %s (+%d)"
                           % (sentido, ", ".join(mudou[:3]), pesos.get("negacao", 20)))

    if tipo in ("ALTERADO", "INCLUIDO", "REMOVIDO"):
        if tipo == "ALTERADO":
            mudado = sum(len(t) for op, t in oc.get("diff", []) if op in "-+")
        else:
            mudado = len(oc.get("texto_modelo", "") or oc.get("texto_devolucao", ""))
        tamanho = min(float(pesos.get("tamanho_max", 15)), mudado / 25.0)
        if tamanho >= 1:
            risco += tamanho
            fatores.append("extens\u00e3o da mudan\u00e7a: %d caracteres (+%d)"
                           % (mudado, tamanho))

    return int(round(max(0, min(100, risco)))), fatores, termos_hit


def _clausula_critica(titulo, termos):
    alvo = sem_acento(titulo or "").lower()
    return any(sem_acento(t).lower() in alvo for t in termos)


# ----------------------------------------------------------------------------
# COMPARACAO
# ----------------------------------------------------------------------------

def _nova_ocorrencia(tipo, a, b, motivo, razao):
    return {
        "tipo": tipo,
        "clausula": (b or a)["clausula"] if b is not None else a["clausula"],
        "ordem_clausula": (b or a).get("ordem_clausula", 0),
        "rotulo": (b or a).get("rotulo", ""),
        "pos_modelo": a["pos"] if a is not None else "",
        "pos_devolucao": b["pos"] if b is not None else "",
        "similaridade": round(razao * 100, 1),
        "texto_modelo": a["texto"] if a is not None else "",
        "texto_devolucao": b["texto"] if b is not None else "",
        "tabela": bool((a or b).get("tipo") == "linha_tabela"),
        "motivo_pareamento": motivo,
        "campos": [],
        "pendentes": [],
        "tambem_movido": motivo == "distante",
    }


def _classificar_par(a, b, motivo, razao, cfg, regexes):
    if a["norm"] == b["norm"]:
        if a.get("placeholders"):
            oc = _nova_ocorrencia("PENDENTE", a, b, motivo, 1.0)
            oc["pendentes"] = list(a["placeholders"])
            oc["tambem_movido"] = motivo in ("movido", "distante")
            return oc
        if motivo == "igual":
            return None
        return _nova_ocorrencia("MOVIDO", a, b, motivo, 1.0)

    casado = casar_modelo(a, b, regexes)
    if casado is not None:
        tipo = "PENDENTE" if casado["pendentes"] else "PREENCHIDO"
        oc = _nova_ocorrencia(tipo, a, b, motivo, razao or _razao(a["norm"], b["norm"]))
        oc["campos"] = casado["campos"]
        oc["pendentes"] = casado["pendentes"]
        return oc
    if chave_formatacao(a["texto"]) == chave_formatacao(b["texto"]):
        return _nova_ocorrencia("FORMATACAO", a, b, motivo, razao or 1.0)
    oc = _nova_ocorrencia("ALTERADO", a, b, motivo, razao or _razao(a["norm"], b["norm"]))
    # campos do modelo que continuam sem preencher mesmo com a redacao alterada
    oc["pendentes"] = [p for p in a.get("placeholders", []) if p in b["norm"]]
    return oc


def comparar(A, B, cfg, alinhamento=None):
    """A = blocos do modelo, B = blocos da devolucao. Lista de ocorrencias.
    Se `alinhamento` (lista) for informado, recebe as linhas [pos_modelo, pos_devolucao]
    de todo o documento na ordem da devolucao (None = sem par), usado na visao
    lado a lado e no documento completo."""
    regexes = _regexes_placeholder(cfg)
    seq_a = [b["norm"] for b in A]
    seq_b = [b["norm"] for b in B]
    sm = difflib.SequenceMatcher(None, seq_a, seq_b, autojunk=False)

    pares, sobra_a, sobra_b = [], [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            pares.extend((i1 + k, j1 + k, "igual", 1.0) for k in range(i2 - i1))
            continue
        p, ra, rb = _parear_faixa(list(range(i1, i2)), list(range(j1, j2)), A, B,
                                  cfg, regexes)
        pares.extend(p)
        sobra_a.extend(ra)
        sobra_b.extend(rb)

    if cfg.get("detectar_blocos_movidos", True) and sobra_a and sobra_b:
        p, sobra_a, sobra_b = _parear_faixa(sobra_a, sobra_b, A, B, cfg, regexes,
                                            distante=True)
        pares.extend(p)

    # posicao equivalente na devolucao para ordenar removidos no lugar certo
    a_para_b = {i: j for i, j, _, _ in pares}

    def posicao(i, j):
        if j is not None:
            return (j, 0)
        anteriores = [a_para_b[k] for k in a_para_b if k < i]
        return ((max(anteriores) if anteriores else -1) + 0.5, i)

    ocorrencias = []
    for i, j, motivo, razao in pares:
        oc = _classificar_par(A[i], B[j], motivo, razao, cfg, regexes)
        if oc:
            oc["_ordem"] = posicao(i, j)
            ocorrencias.append(oc)
    for i in sobra_a:
        oc = _nova_ocorrencia("REMOVIDO", A[i], None, "sem_par", 0.0)
        oc["_ordem"] = posicao(i, None)
        ocorrencias.append(oc)
    for j in sobra_b:
        oc = _nova_ocorrencia("INCLUIDO", None, B[j], "sem_par", 0.0)
        oc["_ordem"] = posicao(None, j)
        ocorrencias.append(oc)
    ocorrencias.sort(key=lambda o: o.pop("_ordem"))

    if alinhamento is not None:
        linhas = [(posicao(i, j), [A[i]["pos"], B[j]["pos"]]) for i, j, _, _ in pares]
        linhas += [(posicao(i, None), [A[i]["pos"], None]) for i in sobra_a]
        linhas += [(posicao(None, j), [None, B[j]["pos"]]) for j in sobra_b]
        alinhamento.extend(l for _, l in sorted(linhas, key=lambda x: x[0]))

    termos = _termos_com_peso(cfg)
    ocultas = set(cfg.get("categorias_ocultas") or [])
    if cfg.get("ignorar_diferencas_apenas_formatacao", True):
        ocultas.add("FORMATACAO")
    for n, oc in enumerate(ocorrencias, 1):
        oc["id"] = "oc-%d" % n
        oc["chave"] = hashlib.sha1(("%s|%s|%s" % (
            oc["tipo"], normalizar(oc["texto_modelo"]), normalizar(oc["texto_devolucao"])
        )).encode("utf-8")).hexdigest()[:12]
        if oc["texto_modelo"] and oc["texto_devolucao"]:
            oc["diff"] = diff_ops(oc["texto_modelo"], oc["texto_devolucao"])
        elif oc["texto_modelo"]:
            oc["diff"] = [["-", oc["texto_modelo"]]]
        else:
            oc["diff"] = [["+", oc["texto_devolucao"]]]
        oc["html_modelo"], oc["html_devolucao"], oc["marcado"] = html_de_ops(oc["diff"])
        if oc["tipo"] == "MOVIDO":
            oc["html_devolucao"] = '<span class="mov">%s</span>' % html.escape(
                oc["texto_devolucao"])
            oc["marcado"] = "[movido] %s" % oc["texto_devolucao"]
        critica = _clausula_critica(oc["clausula"], termos)
        oc["clausula_critica"] = critica
        oc["risco"], oc["fatores"], oc["termos"] = avaliar_risco(oc, cfg, critica)
        oc["selo"] = _selo(oc["risco"], cfg)
        oc["critico"] = _eh_critico(oc, critica, cfg)
        oc["oculto"] = oc["tipo"] in ocultas
    return ocorrencias


def chave_texto(texto):
    """Identidade de um trecho para a biblioteca de clausulas aprovadas."""
    return hashlib.sha1(chave_formatacao(texto).encode("utf-8")).hexdigest()[:16]


def marcar_aprovadas_antes(ocorrencias, biblioteca):
    """biblioteca = {chave_texto: {"data", "fornecedor", ...}} (ja aprovadas)."""
    if not biblioteca:
        return
    for oc in ocorrencias:
        if oc["tipo"] not in ("ALTERADO", "INCLUIDO", "MOVIDO") or not oc["texto_devolucao"]:
            continue
        info = biblioteca.get(chave_texto(oc["texto_devolucao"]))
        if info:
            oc["aprovado_antes"] = {"data": info.get("data", ""),
                                    "fornecedor": info.get("fornecedor", "")}


def _eh_critico(oc, clausula_critica, cfg):
    """Criterio de 'critico' (compativel com v1.1) + risco alto."""
    if oc["tipo"] in ("PREENCHIDO", "FORMATACAO", "PENDENTE"):
        return False
    if oc["selo"] == "Alto":
        return True
    if oc["termos"] or clausula_critica:
        return True
    if oc["tipo"] == "ALTERADO":
        return sorted(valores_do_texto(oc["texto_modelo"])) != sorted(
            valores_do_texto(oc["texto_devolucao"]))
    return False


# ----------------------------------------------------------------------------
# FICHA DO CONTRATO
# ----------------------------------------------------------------------------

RE_DATA = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b")
RE_PERIODO = re.compile(r"(\d{1,2}/\d{1,2}/\d{4})\s*(?:a|at[e\u00e9]|-|\u2013|e)\s*"
                        r"(\d{1,2}/\d{1,2}/\d{4})", re.IGNORECASE)
RE_DINHEIRO = re.compile(r"R\$\s?\d[\d.]*(?:,\d{1,2})?")
RE_PERCENTUAL = re.compile(r"(\d+(?:[.,]\d+)?)\s?%")
RE_PRAZO = re.compile(r"(\d+)\s*(?:\([^)]{1,30}\)\s*)?"
                      r"(dias?\s+[u\u00fa]teis|dias?\s+[u\u00fa]til|dias?\s+corridos|dias?|horas?|h\b|"
                      r"meses|m[e\u00ea]s|anos?)",
                      re.IGNORECASE)
RE_FORO = re.compile(r"foro\s+(?:da\s+comarca\s+|central\s+)?(?:de|da|do)\s+"
                     r"(.+?)(?:\s*[,.;]|\s+para\b|\s+com\b|$)", re.IGNORECASE)


def _data(texto):
    m = RE_DATA.search(texto or "")
    if not m:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def _numero_br(texto):
    t = re.sub(r"[^\d,.]", "", texto or "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _para_horas(numero, unidade):
    u = sem_acento(unidade).lower()
    if u.startswith("h"):
        return float(numero)
    if u.startswith("dia"):
        return float(numero) * (8 if "util" in u or "uteis" in u else 24)
    if u.startswith("mes"):
        return float(numero) * 24 * 30
    if u.startswith("ano"):
        return float(numero) * 24 * 365
    return None


def _contexto(texto, inicio, fim, largura=70):
    a = max(0, inicio - largura)
    b = min(len(texto), fim + 25)
    return ("\u2026" if a else "") + texto[a:b].strip() + ("\u2026" if b < len(texto) else "")


def _rotulo_valor(bloco):
    """Linhas de tabela no formato rotulo | valor."""
    if bloco.get("tipo") == "linha_tabela":
        cel = [c for c in bloco.get("celulas", []) if c]
        if len(cel) == 2:
            return sem_acento(cel[0]).lower(), cel[1]
    m = re.match(r"^\s*([^:]{3,40}):\s*(.+)$", bloco["texto"])
    if m and bloco.get("tipo") != "linha_tabela":
        return sem_acento(m.group(1)).lower(), m.group(2)
    return None, None


def extrair_ficha(B, cfg, fornecedor, cnpj):
    """Campos principais do contrato a partir da devolucao."""
    regexes = _regexes_placeholder(cfg)
    ident = identidade_mod.mesclar(cfg.get("identidade"))
    ficha = {
        "partes": {"contratante": ident["organizacao_contratante"],
                   "contratada": fornecedor, "cnpj_contratada": cnpj},
        "numero_contrato": None, "responsaveis": [], "cnpjs": [],
        "vigencia": {"inicio": None, "fim": None, "meses": None, "valida": None},
        "valores": [], "percentuais": [], "multas": [], "prazos": [], "slas": [],
        "pagamento_dias": None, "foro": None, "alertas": [],
    }
    vistos_cnpj = set()
    regex_numero = None
    try:
        regex_numero = re.compile(cfg.get("regex_numero_contrato")
                                  or CONFIG_PADRAO["regex_numero_contrato"])
    except re.error:
        pass

    for b in B:
        texto = b["texto"]
        rotulo, valor = _rotulo_valor(b)
        tem_campo = bool(encontrar_placeholders(texto, regexes))

        for m in RE_CNPJ.finditer(texto):
            numero = m.group(1)
            if numero in vistos_cnpj:
                continue
            vistos_cnpj.add(numero)
            valido = validar_cnpj(numero)
            ficha["cnpjs"].append({"numero": numero, "valido": valido,
                                   "contexto": _contexto(texto, m.start(), m.end())})
            if not valido:
                ficha["alertas"].append("CNPJ %s com d\u00edgitos verificadores inv\u00e1lidos."
                                        % numero)

        if rotulo:
            if rotulo.startswith("responsavel") and valor and not tem_campo:
                ficha["responsaveis"].append({"papel": b["celulas"][0]
                                              if b.get("celulas") else rotulo,
                                              "nome": valor})
            if ("contrato" in rotulo and ("principal" in rotulo or "numero" in rotulo
                                          or "n" == rotulo.split()[-1])) and valor:
                if not encontrar_placeholders(valor, regexes):
                    ficha["numero_contrato"] = ficha["numero_contrato"] or valor

        if not ficha["numero_contrato"] and regex_numero:
            m = regex_numero.search(texto)
            if m:
                ficha["numero_contrato"] = m.group(0)

        # vigencia: periodo explicito
        clausula_vig = "vig" in sem_acento(b["clausula"]).lower()
        m = RE_PERIODO.search(texto)
        if m and not ficha["vigencia"]["inicio"] and (
                clausula_vig or (rotulo and ("periodo" in rotulo or "vigencia" in rotulo))
                or "vig" in sem_acento(texto).lower()):
            ficha["vigencia"]["inicio"] = m.group(1)
            ficha["vigencia"]["fim"] = m.group(2)
        if ficha["vigencia"]["meses"] is None and (clausula_vig or "vigor" in texto.lower()):
            mm = re.search(r"(\d+)\s*(?:\([^)]{1,20}\)\s*)?(meses|anos?)", texto,
                           re.IGNORECASE)
            if mm:
                n = int(mm.group(1))
                ficha["vigencia"]["meses"] = n * (12 if mm.group(2).lower().startswith("ano")
                                                  else 1)

        for m in RE_DINHEIRO.finditer(texto):
            ficha["valores"].append({"valor": m.group(0).replace("R$ ", "R$ "),
                                     "numero": _numero_br(m.group(0)),
                                     "rotulo": (b["celulas"][0] if rotulo and b.get("celulas")
                                                else ""),
                                     "contexto": _contexto(texto, m.start(), m.end())})

        for m in RE_PERCENTUAL.finditer(texto):
            ctx = _contexto(texto, m.start(), m.end())
            item = {"valor": m.group(0), "numero": _numero_br(m.group(1)), "contexto": ctx}
            ficha["percentuais"].append(item)
            janela = sem_acento(texto[max(0, m.start() - 90):m.end()]).lower()
            if "multa" in janela or "penalidade" in janela:
                ficha["multas"].append(item)

        for m in RE_PRAZO.finditer(texto):
            ctx = _contexto(texto, m.start(), m.end())
            horas = _para_horas(int(m.group(1)), m.group(2))
            janela = sem_acento(texto[max(0, m.start() - 110):m.end()]).lower()
            item = {"valor": m.group(0), "horas": horas, "contexto": ctx}
            ficha["prazos"].append(item)
            if ("pagamento" in janela or "pagar" in janela or "pago" in janela) and \
                    "dia" in m.group(2).lower() and ficha["pagamento_dias"] is None:
                ficha["pagamento_dias"] = int(m.group(1))

        titulo_clausula = sem_acento(b["clausula"]).lower()
        if b.get("tipo") == "linha_tabela" and ("sla" in titulo_clausula
                                                or "indicador" in titulo_clausula):
            cel = [c for c in b.get("celulas", []) if c]
            if cel and not b.get("linha") == 0 and not tem_campo:
                horas = None
                for c in cel[1:]:
                    mh = RE_PRAZO.search(c)
                    if mh:
                        horas = _para_horas(int(mh.group(1)), mh.group(2))
                        break
                    if re.search(r"\bdia\s+[u\u00fa]til\b", c, re.IGNORECASE):  # 1 dia util = 8h
                        horas = 8.0
                        break
                ficha["slas"].append({"indicador": cel[0], "linha": " | ".join(cel),
                                      "horas": horas})

        if ficha["foro"] is None:
            m = RE_FORO.search(texto)
            if m:
                cidade = m.group(1).strip()
                pendente = bool(encontrar_placeholders(cidade, regexes))
                ficha["foro"] = {"cidade": None if pendente else cidade,
                                 "pendente": pendente,
                                 "contexto": _contexto(texto, m.start(), m.end())}

    ini, fim = _data(ficha["vigencia"]["inicio"]), _data(ficha["vigencia"]["fim"])
    if ini and fim:
        ficha["vigencia"]["valida"] = ini < fim
        if ficha["vigencia"]["meses"] is None:
            ficha["vigencia"]["meses"] = max(0, round((fim - ini).days / 30.44))
        if not ini < fim:
            ficha["alertas"].append("Vig\u00eancia com in\u00edcio (%s) posterior ao fim (%s)."
                                    % (ficha["vigencia"]["inicio"], ficha["vigencia"]["fim"]))
    if cnpj and not validar_cnpj(cnpj):
        ficha["alertas"].append("CNPJ da contratada (%s) inv\u00e1lido." % cnpj)
    if ficha["foro"] and ficha["foro"]["pendente"]:
        ficha["alertas"].append("Foro n\u00e3o preenchido.")

    multas = [m["numero"] for m in ficha["multas"] if m["numero"] is not None]
    horas_sla = [s["horas"] for s in ficha["slas"] if s["horas"] is not None]
    ficha["indicadores"] = {
        "multa_percentual_max": max(multas) if multas else None,
        "juros_percentual_max": max([p["numero"] for p in ficha["percentuais"]
                                     if p["numero"] is not None
                                     and "juro" in sem_acento(p["contexto"]).lower()]
                                    or [None], key=lambda x: -1 if x is None else x),
        "pagamento_dias": ficha["pagamento_dias"],
        "vigencia_meses": ficha["vigencia"]["meses"],
        "foro_cidade": (ficha["foro"] or {}).get("cidade"),
        "sla_horas_min": min(horas_sla) if horas_sla else None,
        "sla_horas_max": max(horas_sla) if horas_sla else None,
        "valor_total_reais": sum(v["numero"] for v in ficha["valores"]
                                 if v["numero"] is not None) or None,
    }
    return ficha


# ----------------------------------------------------------------------------
# PLAYBOOK DE ACEITACAO
# ----------------------------------------------------------------------------

CAMPOS_PLAYBOOK = {
    "multa_percentual_max": "Multa (maior %)",
    "juros_percentual_max": "Juros (maior %)",
    "pagamento_dias": "Prazo de pagamento (dias)",
    "vigencia_meses": "Vig\u00eancia (meses)",
    "foro_cidade": "Foro (cidade)",
    "sla_horas_min": "SLA mais curto (horas)",
    "sla_horas_max": "SLA mais longo (horas)",
    "valor_total_reais": "Soma dos valores (R$)",
}
OPERADORES_PLAYBOOK = ("<=", ">=", "<", ">", "=", "!=")


def padrao_valor_playbook(campo, valor):
    """Regex do valor encontrado COM a unidade do campo (evita casar '10' em datas)."""
    if isinstance(valor, str):
        return re.compile(re.escape(valor), re.IGNORECASE) if valor.strip() else None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    numero = re.escape(str(valor)).replace("\\.", "[.,]")
    antes = r"(?<![\d/.,])"
    if campo.endswith("_percentual_max"):
        return re.compile(antes + numero + r"(?:[.,]0+)?\s?%")
    if campo == "pagamento_dias":
        return re.compile(antes + numero + r"\s*(?:\([^)]{1,30}\)\s*)?dias?", re.I)
    if campo == "vigencia_meses":
        return re.compile(antes + numero + r"\s*(?:\([^)]{1,30}\)\s*)?mes", re.I)
    if campo.startswith("sla_horas"):
        return re.compile(antes + numero + r"\s*(?:h\b|hora)", re.I)
    return None


def avaliar_playbook(ficha, regras):
    """Confere os indicadores da ficha contra as regras de aceitacao.
    regra = {"id","rotulo","campo","operador","valor","ativo"}"""
    resultado = []
    indicadores = ficha.get("indicadores", {})
    for regra in regras or []:
        if not regra.get("ativo", True):
            continue
        campo, op, esperado = regra.get("campo"), regra.get("operador"), regra.get("valor")
        rotulo = regra.get("rotulo") or "%s %s %s" % (CAMPOS_PLAYBOOK.get(campo, campo),
                                                      op, esperado)
        encontrado = indicadores.get(campo)
        status = "sem_dado"
        if encontrado is not None and op in OPERADORES_PLAYBOOK:
            if isinstance(encontrado, str) or isinstance(esperado, str):
                a = sem_acento(str(encontrado)).strip().lower()
                b = sem_acento(str(esperado)).strip().lower()
                ok = (a == b or b in a) if op == "=" else (a != b if op == "!=" else None)
            else:
                try:
                    e = float(esperado)
                    ok = {"<=": encontrado <= e, ">=": encontrado >= e, "<": encontrado < e,
                          ">": encontrado > e, "=": encontrado == e,
                          "!=": encontrado != e}[op]
                except (TypeError, ValueError):
                    ok = None
            if ok is not None:
                status = "ok" if ok else "fora"
        resultado.append({"id": regra.get("id") or campo, "rotulo": rotulo,
                          "campo": campo, "operador": op, "esperado": esperado,
                          "encontrado": encontrado, "status": status})
    return resultado


# ----------------------------------------------------------------------------
# ANALISE COMPLETA
# ----------------------------------------------------------------------------

def montar_resumo(ocorrencias):
    """resumo[tipo] = {total, criticos} (inclui chaves v1.1) e contagem por clausula."""
    resumo = {t: {"total": 0, "criticos": 0} for t in TIPOS + ("TOTAL",)}
    por_clausula = {}
    for a in ocorrencias:
        tipo = a["tipo"]
        resumo[tipo]["total"] += 1
        resumo["TOTAL"]["total"] += 1
        if a.get("critico"):
            resumo[tipo]["criticos"] += 1
            resumo["TOTAL"]["criticos"] += 1
        cont = por_clausula.setdefault(a["clausula"], dict(
            {t: 0 for t in TIPOS}, ordem=a.get("ordem_clausula", 0), risco_max=0,
            criticos=0))
        cont[tipo] += 1
        cont["risco_max"] = max(cont["risco_max"], a.get("risco", 0))
        cont["criticos"] += 1 if a.get("critico") else 0
    ordenado = sorted(por_clausula.items(), key=lambda kv: kv[1]["ordem"])
    return resumo, ordenado


def risco_geral(ocorrencias, cfg):
    visiveis = [o["risco"] for o in ocorrencias if not o.get("oculto")]
    if not visiveis:
        return 0
    top = sorted(visiveis, reverse=True)[:5]
    altos = sum(1 for o in ocorrencias if o.get("selo") == "Alto")
    valor = 0.5 * top[0] + 0.3 * (sum(top) / len(top)) + min(20, altos * 4)
    return int(round(min(100, valor)))


def _clausulas_documento(blocos):
    vistas, saida = set(), []
    for b in blocos:
        if b["clausula"] in vistas:
            continue
        vistas.add(b["clausula"])
        saida.append({"titulo": b["clausula"], "ordem": b.get("ordem_clausula", 0),
                      "pos": b["pos"]})
    return saida


def analisar(caminho_modelo, caminho_devolucao, cfg, log=None, agora=None):
    """Analise completa, sem gravar nada. Devolve o dicionario do resultado.json."""
    agora = agora or datetime.now()
    blocos_dev = montar_blocos(ler_documento_estruturado(caminho_devolucao), cfg)
    if not blocos_dev:
        raise ValueError("Nenhum texto extraido da devolucao (PDF de imagem? precisa OCR)")
    blocos_mod = montar_blocos(ler_documento_estruturado(caminho_modelo), cfg)

    fornecedor, cnpj = detectar_fornecedor(caminho_devolucao, cfg, blocos_dev)
    alinhamento = []
    ocorrencias = comparar(blocos_mod, blocos_dev, cfg, alinhamento)
    marcar_aprovadas_antes(ocorrencias, cfg.get("biblioteca_aprovada"))
    resumo, por_clausula = montar_resumo(ocorrencias)
    ficha = extrair_ficha(blocos_dev, cfg, fornecedor, cnpj)
    playbook = avaliar_playbook(ficha, cfg.get("playbook"))

    # ocorrencias ligadas a regras fora da politica (valor + unidade no texto)
    for regra in playbook:
        if regra["status"] != "fora" or regra["encontrado"] is None:
            continue
        padrao = padrao_valor_playbook(regra["campo"], regra["encontrado"])
        if not padrao:
            continue
        for oc in ocorrencias:
            if oc["tipo"] != "FORMATACAO" and padrao.search(oc["texto_devolucao"] or ""):
                oc.setdefault("fora_politica", []).append(regra["rotulo"])

    pendencias = []
    for oc in ocorrencias:
        for p in oc.get("pendentes", []):
            pendencias.append({"placeholder": p, "clausula": oc["clausula"],
                               "ocorrencia": oc["id"], "pos_devolucao": oc["pos_devolucao"]})

    sim_global = difflib.SequenceMatcher(
        None, [b["norm"] for b in blocos_mod], [b["norm"] for b in blocos_dev],
        autojunk=False).ratio() * 100
    geral = risco_geral(ocorrencias, cfg)

    meta = {
        "data_hora": agora.strftime("%d/%m/%Y %H:%M:%S"),
        "data_hora_iso": agora.isoformat(timespec="seconds"),
        "fornecedor": fornecedor,
        "cnpj": cnpj,
        "arquivo_modelo": os.path.basename(caminho_modelo),
        "arquivo_devolucao": os.path.basename(caminho_devolucao),
        "qtd_modelo": len(blocos_mod),
        "qtd_devolucao": len(blocos_dev),
        "similaridade_global": sim_global,
        "versao_motor": VERSAO,
        "risco_geral": geral,
        "selo_risco": _selo(geral, cfg),
        "clausulas_detectadas": len({b["clausula"] for b in blocos_dev}),
        "devolucao_incompleta": bool(pendencias),
        "fora_da_politica": sum(1 for r in playbook if r["status"] == "fora"),
        "aprovadas_antes": sum(1 for o in ocorrencias if o.get("aprovado_antes")),
        "visao_padrao": cfg.get("visao_padrao", "inline"),
        "categorias_ocultas": sorted({o["tipo"] for o in ocorrencias if o["oculto"]}
                                     | set(cfg.get("categorias_ocultas") or [])),
    }
    try:
        meta["hash_entrada"] = hash_entrada(caminho_modelo, caminho_devolucao, cfg)
    except OSError:
        meta["hash_entrada"] = ""

    resultado = {
        "versao_formato": VERSAO_FORMATO_RESULTADO,
        "meta": meta,
        "identidade": identidade_mod.mesclar(cfg.get("identidade")),
        "ficha": ficha,
        "playbook": playbook,
        "resumo": resumo,
        "por_clausula": [dict(v, clausula=k) for k, v in por_clausula],
        "clausulas": {"modelo": _clausulas_documento(blocos_mod),
                      "devolucao": _clausulas_documento(blocos_dev)},
        "pendencias": pendencias,
        "ocorrencias": ocorrencias,
        "config_relevante": config_relevante(cfg),
    }
    if cfg.get("incluir_texto_integral", True):
        ocs_por_pos = {}
        for oc in ocorrencias:
            if oc["pos_devolucao"] != "":
                ocs_por_pos.setdefault(oc["pos_devolucao"], []).append(oc["id"])
        resultado["documento"] = {
            "modelo": [{"pos": b["pos"], "texto": b["texto"], "clausula": b["clausula"],
                        "tipo": b.get("tipo", "paragrafo"), "titulo": b.get("titulo", False)}
                       for b in blocos_mod],
            "devolucao": [{"pos": b["pos"], "texto": b["texto"], "clausula": b["clausula"],
                           "tipo": b.get("tipo", "paragrafo"),
                           "titulo": b.get("titulo", False),
                           "ocorrencias": ocs_por_pos.get(b["pos"], [])}
                          for b in blocos_dev],
        }
        resultado["alinhamento"] = alinhamento
    if log:
        log.info("   Fornecedor: %s%s", fornecedor, (" | CNPJ %s" % cnpj) if cnpj else "")
        log.info("   Ocorrencias: %s | risco geral %d (%s)",
                 ", ".join("%d %s" % (resumo[t]["total"], t.lower()) for t in TIPOS
                           if resumo[t]["total"]), geral, meta["selo_risco"])
    return resultado


# ----------------------------------------------------------------------------
# ORQUESTRACAO
# ----------------------------------------------------------------------------

def escrever_com_retry(caminho, conteudo, log, binario=False):
    for tentativa in range(3):
        try:
            tmp = caminho + ".tmp"
            if binario:
                with open(tmp, "wb") as f:
                    f.write(conteudo)
            else:
                with open(tmp, "w", encoding="utf-8") as f:
                    f.write(conteudo)
            os.replace(tmp, caminho)
            return True
        except PermissionError:
            log.warning("Arquivo em uso (%s). Nova tentativa em 3s...", caminho)
            time.sleep(3)
    log.error("Nao foi possivel gravar: %s", caminho)
    return False


def preparar_ambiente(cfg):
    for chave in ("pasta_modelo", "pasta_devolucoes", "pasta_saida", "pasta_logs"):
        os.makedirs(cfg[chave], exist_ok=True)


def configurar_log(cfg):
    log = logging.getLogger("contr")
    log.setLevel(logging.INFO)
    log.handlers = []
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s",
                            "%Y-%m-%d %H:%M:%S")
    arquivo = os.path.join(cfg["pasta_logs"],
                           "comparar_contratos_%s.log" % datetime.now().strftime("%Y%m%d"))
    fh = logging.FileHandler(arquivo, encoding="utf-8")
    fh.setFormatter(fmt)
    ch = logging.StreamHandler(sys.stdout)
    ch.setFormatter(logging.Formatter("%(message)s"))
    log.addHandler(fh)
    log.addHandler(ch)
    return log


def listar_arquivos(pasta, extensoes, recursivo=True):
    encontrados = []
    for raiz, _, arquivos in os.walk(pasta):
        for nome in sorted(arquivos):
            if nome.startswith("~$") or nome.startswith("."):
                continue
            if os.path.splitext(nome)[1].lower() in extensoes:
                encontrados.append(os.path.join(raiz, nome))
        if not recursivo:
            break
    return sorted(encontrados)


def escolher_modelo(caminho_devolucao, modelos, cfg, log):
    if cfg.get("modelo_fixo"):
        alvo = cfg["modelo_fixo"]
        if not os.path.isabs(alvo):
            alvo = os.path.join(cfg["pasta_modelo"], alvo)
        if os.path.exists(alvo):
            return alvo
        log.warning("modelo_fixo nao encontrado: %s", alvo)

    base = os.path.basename(caminho_devolucao)
    for chave, valor in (cfg.get("mapa_modelos") or {}).items():
        if sem_acento(chave).lower() in sem_acento(base).lower():
            alvo = valor if os.path.isabs(valor) else os.path.join(cfg["pasta_modelo"], valor)
            if os.path.exists(alvo):
                return alvo

    if len(modelos) == 1:
        return modelos[0]

    alvo_norm = sem_acento(os.path.splitext(base)[0]).lower()
    melhor, melhor_r = modelos[0], -1.0
    for m in modelos:
        nome = sem_acento(os.path.splitext(os.path.basename(m))[0]).lower()
        r = difflib.SequenceMatcher(None, alvo_norm, nome).ratio()
        if r > melhor_r:
            melhor, melhor_r = m, r
    log.info("   modelo escolhido por similaridade de nome (%.0f%%): %s",
             melhor_r * 100, os.path.basename(melhor))
    return melhor


def criar_pasta_saida(cfg, fornecedor, agora):
    partes = [sanitizar_nome_pasta(fornecedor),
              agora.strftime(cfg.get("formato_data_pasta", "%d.%m.%Y"))]
    nome = "-".join(partes)
    if cfg.get("incluir_hora_na_pasta", True):
        nome += "_" + agora.strftime(cfg.get("formato_hora_pasta", "%Hh%M"))
    destino = os.path.join(cfg["pasta_saida"], nome)
    sufixo = 2
    while os.path.exists(destino):
        destino = os.path.join(cfg["pasta_saida"], "%s_%d" % (nome, sufixo))
        sufixo += 1
    os.makedirs(destino, exist_ok=True)
    return destino


def registrar_historico(cfg, linha):
    caminho = os.path.join(cfg["pasta_saida"], cfg.get("historico_csv",
                                                       "_historico_comparacoes.csv"))
    novo = not os.path.exists(caminho)
    with open(caminho, "a", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        if novo:
            w.writerow(["data_hora", "fornecedor", "cnpj", "arquivo_modelo",
                        "arquivo_devolucao", "similaridade_global", "alterados",
                        "incluidos", "removidos", "movidos", "criticos", "pasta_saida"])
        w.writerow(linha)


def gravar_saidas(resultado, destino, cfg, log):
    """Grava resultado.json e os relatorios configurados na pasta destino."""
    import relatorios
    arquivos = []
    if cfg.get("gerar_json", True):
        escrever_com_retry(os.path.join(destino, "resultado.json"),
                           json.dumps(resultado, ensure_ascii=False, indent=1), log)
        arquivos.append("resultado.json")
    if cfg.get("gerar_html", True):
        escrever_com_retry(os.path.join(destino, "comparativo.html"),
                           relatorios.montar_html_interativo(resultado, cfg), log)
        arquivos.append("comparativo.html")
    if cfg.get("gerar_pdf", True):
        if relatorios.gravar_pdf(relatorios.montar_html_pdf(resultado, cfg),
                                 os.path.join(destino, "comparativo.pdf"), log):
            arquivos.append("comparativo.pdf")
    if cfg.get("gerar_txt", True):
        escrever_com_retry(os.path.join(destino, "resumo_alteracoes.txt"),
                           relatorios.montar_txt(resultado, cfg), log)
        arquivos.append("resumo_alteracoes.txt")
    if cfg.get("gerar_xlsx", True):
        if relatorios.gravar_xlsx(os.path.join(destino, "alteracoes.xlsx"), resultado,
                                  cfg, log):
            arquivos.append("alteracoes.xlsx")
    return arquivos


def processar(caminho_devolucao, modelos, cfg, log):
    """Compara uma devolucao e grava as saidas. Retorna (meta, resumo, destino);
    meta['resultado'] traz o dicionario completo."""
    agora = datetime.now()
    log.info("-" * 72)
    raiz = cfg.get("pasta_devolucoes") or os.path.dirname(caminho_devolucao)
    log.info("Devolucao: %s", os.path.relpath(caminho_devolucao, raiz))

    caminho_modelo = escolher_modelo(caminho_devolucao, modelos, cfg, log)
    log.info("   Modelo: %s", os.path.basename(caminho_modelo))
    resultado = analisar(caminho_modelo, caminho_devolucao, cfg, log, agora)
    meta, resumo = resultado["meta"], resultado["resumo"]

    if not cfg.get("gravar_saida", True):
        log.info("   [gravar_saida=false] nada gravado.")
        return dict(meta, resultado=resultado), resumo, None

    destino = criar_pasta_saida(cfg, meta["fornecedor"], agora)
    log.info("   Saida: %s", destino)
    resultado["meta"]["pasta_saida"] = os.path.basename(destino)
    gravar_saidas(resultado, destino, cfg, log)

    if cfg.get("copiar_arquivos_origem", False):
        import shutil
        origem = os.path.join(destino, "_originais")
        os.makedirs(origem, exist_ok=True)
        shutil.copy2(caminho_modelo, origem)
        shutil.copy2(caminho_devolucao, origem)

    registrar_historico(cfg, [
        meta["data_hora"], meta["fornecedor"], meta["cnpj"], meta["arquivo_modelo"],
        meta["arquivo_devolucao"], "%.1f" % meta["similaridade_global"],
        resumo["ALTERADO"]["total"], resumo["INCLUIDO"]["total"],
        resumo["REMOVIDO"]["total"], resumo["MOVIDO"]["total"],
        resumo["TOTAL"]["criticos"], destino])

    return dict(meta, resultado=resultado), resumo, destino


def main():
    cfg = carregar_config()
    preparar_ambiente(cfg)
    log = configurar_log(cfg)

    log.info("=" * 72)
    log.info("comparar_contratos v%s | %s", VERSAO,
             datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
    if cfg.get("modo_teste", True):
        log.info("MODO_TESTE ativo - limite de %d arquivo(s)",
                 cfg.get("limite_arquivos_teste", 1))

    extensoes = [e.lower() for e in cfg.get("extensoes_aceitas", [])]
    modelos = listar_arquivos(cfg["pasta_modelo"], extensoes)
    if not modelos:
        log.error("Nenhum modelo padrao em: %s", cfg["pasta_modelo"])
        return
    devolucoes = listar_arquivos(cfg["pasta_devolucoes"], extensoes)
    if not devolucoes:
        log.error("Nenhuma devolucao em: %s", cfg["pasta_devolucoes"])
        return
    if cfg.get("modo_teste", True):
        devolucoes = devolucoes[:int(cfg.get("limite_arquivos_teste", 1))]

    log.info("Modelos: %d | Devolucoes a processar: %d", len(modelos), len(devolucoes))

    processados, falhas = [], []
    for caminho in devolucoes:
        try:
            meta, resumo, destino = processar(caminho, modelos, cfg, log)
            processados.append((meta, resumo, destino))
        except Exception as erro:  # noqa: BLE001
            log.exception("FALHA em %s: %s", os.path.basename(caminho), erro)
            falhas.append((os.path.basename(caminho), str(erro)))

    log.info("=" * 72)
    log.info("%-30s %-24s %4s %4s %4s %4s %4s %4s %5s", "ARQUIVO", "FORNECEDOR",
             "PRE", "ALT", "INC", "REM", "MOV", "PEN", "RISCO")
    for meta, resumo, _ in processados:
        log.info("%-30.30s %-24.24s %4d %4d %4d %4d %4d %4d %5d",
                 meta["arquivo_devolucao"], meta["fornecedor"],
                 resumo["PREENCHIDO"]["total"], resumo["ALTERADO"]["total"],
                 resumo["INCLUIDO"]["total"], resumo["REMOVIDO"]["total"],
                 resumo["MOVIDO"]["total"], resumo["PENDENTE"]["total"],
                 meta["risco_geral"])
    if falhas:
        log.info("-" * 72)
        for nome, erro in falhas:
            log.info("FALHA: %s -> %s", nome, erro)
    log.info("Concluido: %d ok, %d falha(s).", len(processados), len(falhas))


if __name__ == "__main__":
    main()
